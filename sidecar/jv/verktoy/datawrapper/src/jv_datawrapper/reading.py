"""Reading .xlsx/.xlsm (openpyxl), .xls (xlrd) and .csv/.txt into plain grids.

Every visible sheet becomes a `Sheet`: a rectangular grid of raw values
(None, str, int, float, bool, date/datetime/time) with the hidden rows and
columns already taken out. The original positions are kept so warnings can
point to cells (A1 notation) without quoting their content.
"""

from __future__ import annotations

import codecs
import csv
import io
import os
import warnings
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path

from jv_common import UserError

MAX_BYTES = 50 * 1024 * 1024
MAX_CELLS = 5_000_000

EXCEL_SUFFIXES = (".xlsx", ".xlsm")
XLS_SUFFIXES = (".xls",)
TEXT_SUFFIXES = (".csv", ".txt")
SUFFIXES = EXCEL_SUFFIXES + XLS_SUFFIXES + TEXT_SUFFIXES

_CFB_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
_ZIP_MAGIC = b"PK\x03\x04"
_ENCRYPTED_MARK = "EncryptionInfo".encode("utf-16-le")

EXCEL_ERRORS = frozenset(
    {"#N/A", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#NULL!", "#GETTING_DATA", "#SPILL!", "#CALC!"}
)


@dataclass
class Sheet:
    name: str
    rows: list[list[object]]
    # Original (0-based) row and column of each row/column in `rows`.
    row_ids: list[int]
    col_ids: list[int]
    # Merged ranges in grid coordinates, inclusive: (r0, c0, r1, c1).
    merges: list[tuple[int, int, int, int]] = field(default_factory=list)
    # Original indices of hidden rows/columns that contained data.
    hidden_rows: list[int] = field(default_factory=list)
    hidden_cols: list[int] = field(default_factory=list)
    # Grid cells with a percent number format in Excel.
    percent: set[tuple[int, int]] = field(default_factory=set)
    # Cells that hold an Excel error value (#DIV/0! ...), grid coordinates.
    errors: set[tuple[int, int]] = field(default_factory=set)

    @property
    def width(self) -> int:
        return len(self.col_ids)

    @property
    def height(self) -> int:
        return len(self.row_ids)


@dataclass
class Workbook:
    kind: str  # "xlsx", "xls" or "csv"
    sheets: list[Sheet]
    hidden_sheets: int = 0
    encoding: str | None = None  # csv only
    delimiter: str | None = None  # csv only


def col_letter(index: int) -> str:
    """0 -> A, 25 -> Z, 26 -> AA."""
    out = ""
    n = index + 1
    while n:
        n, rem = divmod(n - 1, 26)
        out = chr(65 + rem) + out
    return out


def cell_ref(row: int, col: int) -> str:
    return f"{col_letter(col)}{row + 1}"


def is_empty(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


# ---------------------------------------------------------------- the grid


def _build_sheet(
    name: str,
    cells: dict[tuple[int, int], object],
    hidden_rows: set[int],
    hidden_cols: set[int],
    merges: list[tuple[int, int, int, int]],
    percent: set[tuple[int, int]],
) -> Sheet:
    """Grid from sparse cells (original coordinates), hidden rows/columns removed."""
    used = [(r, c) for (r, c), v in cells.items() if not is_empty(v)]
    if not used:
        return Sheet(name, [], [], [])
    max_r = max(r for r, _ in used)
    max_c = max(c for _, c in used)
    if (max_r + 1) * (max_c + 1) > MAX_CELLS:
        raise UserError(
            "Arket er for stort til å klargjøres her (mer enn fem millioner celler). "
            "Kopier tabellen til et nytt ark og prøv igjen."
        )
    data_rows = {r for r, _ in used}
    data_cols = {c for _, c in used}
    row_ids = [r for r in range(max_r + 1) if r not in hidden_rows]
    col_ids = [c for c in range(max_c + 1) if c not in hidden_cols]
    r_pos = {r: i for i, r in enumerate(row_ids)}
    c_pos = {c: i for i, c in enumerate(col_ids)}
    rows: list[list[object]] = [[None] * len(col_ids) for _ in row_ids]
    errors: set[tuple[int, int]] = set()
    for (r, c), v in cells.items():
        if r in r_pos and c in c_pos:
            if isinstance(v, str) and v.strip() in EXCEL_ERRORS:
                errors.add((r_pos[r], c_pos[c]))
                v = None
            rows[r_pos[r]][c_pos[c]] = v
    grid_merges = []
    for r0, c0, r1, c1 in merges:
        if r0 not in r_pos or c0 not in c_pos:
            continue  # the anchor is hidden
        vr = [r_pos[r] for r in range(r0, r1 + 1) if r in r_pos]
        vc = [c_pos[c] for c in range(c0, c1 + 1) if c in c_pos]
        if len(vr) * len(vc) > 1:
            grid_merges.append((vr[0], vc[0], vr[-1], vc[-1]))
    return Sheet(
        name=name,
        rows=rows,
        row_ids=row_ids,
        col_ids=col_ids,
        merges=grid_merges,
        hidden_rows=sorted(r for r in hidden_rows if r in data_rows),
        hidden_cols=sorted(c for c in hidden_cols if c in data_cols),
        percent={(r_pos[r], c_pos[c]) for r, c in percent if r in r_pos and c in c_pos},
        errors=errors,
    )


# ---------------------------------------------------------------- xlsx


def _read_xlsx(path: Path) -> Workbook:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            wb = load_workbook(path, data_only=True, keep_links=False, rich_text=False)
    except (zipfile.BadZipFile, InvalidFileException, KeyError, ValueError, TypeError, EOFError) as exc:
        raise _broken() from exc
    sheets: list[Sheet] = []
    hidden = 0
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            hidden += 1
            continue
        cells: dict[tuple[int, int], object] = {}
        percent: set[tuple[int, int]] = set()
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if v is None:
                    continue
                key = (cell.row - 1, cell.column - 1)
                cells[key] = v
                if isinstance(v, (int, float)) and not isinstance(v, bool) and "%" in (cell.number_format or ""):
                    percent.add(key)
        hidden_rows = {r - 1 for r, dim in ws.row_dimensions.items() if dim.hidden}
        hidden_cols: set[int] = set()
        for dim in ws.column_dimensions.values():
            if dim.hidden and dim.min and dim.max:
                hidden_cols.update(range(dim.min - 1, min(dim.max, 16384)))
        merges = [
            (m.min_row - 1, m.min_col - 1, m.max_row - 1, m.max_col - 1) for m in ws.merged_cells.ranges
        ]
        sheets.append(_build_sheet(ws.title, cells, hidden_rows, hidden_cols, merges, percent))
    return Workbook("xlsx", sheets, hidden)


# ---------------------------------------------------------------- xls


def _read_xls(path: Path) -> Workbook:
    import xlrd

    data = path.read_bytes()
    with open(os.devnull, "w") as devnull:  # xlrd writes warnings to its logfile (stdout by default)
        try:
            book = xlrd.open_workbook(
                file_contents=data, formatting_info=True, logfile=devnull, verbosity=0, on_demand=False
            )
        except xlrd.biffh.XLRDError as exc:
            if "encrypt" in str(exc).lower():
                raise _protected() from exc
            raise _broken() from exc
        except Exception as exc:  # noqa: BLE001 - xlrd raises many kinds on a damaged file
            raise _broken() from exc
    sheets: list[Sheet] = []
    hidden = 0
    for sh in book.sheets():
        if sh.visibility != 0:
            hidden += 1
            continue
        cells: dict[tuple[int, int], object] = {}
        percent: set[tuple[int, int]] = set()
        for r in range(sh.nrows):
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                t = cell.ctype
                if t in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    continue
                if t == xlrd.XL_CELL_DATE:
                    try:
                        v: object = xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)
                    except (ValueError, OverflowError, xlrd.xldate.XLDateError):
                        v = cell.value
                elif t == xlrd.XL_CELL_BOOLEAN:
                    v = bool(cell.value)
                elif t == xlrd.XL_CELL_ERROR:
                    v = xlrd.error_text_from_code.get(cell.value, "#N/A")
                elif t == xlrd.XL_CELL_NUMBER:
                    v = cell.value
                    if _xls_percent(book, sh, r, c):
                        percent.add((r, c))
                else:
                    v = cell.value
                cells[(r, c)] = v
        hidden_rows = {r for r, info in sh.rowinfo_map.items() if info.hidden}
        hidden_cols = {c for c, info in sh.colinfo_map.items() if info.hidden}
        merges = [(rlo, clo, rhi - 1, chi - 1) for rlo, rhi, clo, chi in sh.merged_cells]
        sheets.append(_build_sheet(sh.name, cells, hidden_rows, hidden_cols, merges, percent))
    return Workbook("xls", sheets, hidden)


def _xls_percent(book, sheet, r: int, c: int) -> bool:
    try:
        xf = book.xf_list[sheet.cell_xf_index(r, c)]
        fmt = book.format_map.get(xf.format_key)
        return bool(fmt and "%" in fmt.format_str)
    except (IndexError, AttributeError):
        return False


# ---------------------------------------------------------------- csv


def decode_text(data: bytes) -> tuple[str, str]:
    """Text and encoding name: BOM first, then UTF-8, then Windows-1252."""
    for bom, enc in ((codecs.BOM_UTF8, "utf-8-sig"), (codecs.BOM_UTF16_LE, "utf-16"), (codecs.BOM_UTF16_BE, "utf-16")):
        if data.startswith(bom):
            try:
                return data.decode(enc), "UTF-8 med BOM" if enc == "utf-8-sig" else "UTF-16"
            except UnicodeDecodeError as exc:
                raise _broken() from exc
    try:
        return data.decode("utf-8"), "UTF-8"
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace"), "Windows-1252"


def detect_delimiter(text: str) -> str:
    """`;`, `,` or tab: the one that splits the first lines most consistently.

    Ties go to tab, then `;`, then `,` (a Norwegian `;` file has decimal
    commas in it, so `,` alone is weak evidence).
    """
    lines = [ln for ln in text.splitlines() if ln.strip()][:50]
    best = ("\t", (-1.0, 0))
    for delim in ("\t", ";", ","):
        counts = [len(row) for row in csv.reader(lines, delimiter=delim)]
        if not counts:
            continue
        multi = [n for n in counts if n > 1]
        if not multi:
            continue
        mode = max(set(multi), key=multi.count)
        score = (multi.count(mode) / len(counts), mode)
        if score > best[1]:
            best = (delim, score)
    return best[0] if best[1][0] > 0 else ";"


def _read_csv(path: Path) -> Workbook:
    data = path.read_bytes()
    text, encoding = decode_text(data)
    if not text.strip():
        raise _empty()
    delim = detect_delimiter(text)
    try:
        rows = list(csv.reader(io.StringIO(text, newline=""), delimiter=delim))
    except csv.Error as exc:
        raise _broken() from exc
    cells = {(r, c): v for r, row in enumerate(rows) for c, v in enumerate(row) if v.strip()}
    sheet = _build_sheet("CSV", cells, set(), set(), [], set())
    return Workbook("csv", [sheet], 0, encoding=encoding, delimiter=delim)


# ---------------------------------------------------------------- entry


def _broken() -> UserError:
    return UserError(
        "Filen kan ikke leses. Den ser ut til å være ødelagt, eller den er ikke et regneark "
        "eller en CSV-fil. Åpne den i Excel og lagre den på nytt som .xlsx."
    )


def _protected() -> UserError:
    return UserError(
        "Filen er passordbeskyttet. Åpne den i Excel, fjern passordet og lagre den på nytt."
    )


def _empty() -> UserError:
    return UserError("Filen er tom. Fant ingen data å klargjøre.")


def read_file(path_str: str) -> Workbook:
    path = Path(path_str)
    if not path.is_file():
        raise UserError(f"Finner ikke filen «{path.name}».")
    suffix = path.suffix.lower()
    if suffix not in SUFFIXES:
        raise UserError(
            f"Filtypen «{path.suffix or path.name}» støttes ikke. Bruk .xlsx, .xls eller .csv."
        )
    size = path.stat().st_size
    if size > MAX_BYTES:
        raise UserError("Filen er større enn 50 MB. Kopier tabellen til en mindre fil og prøv igjen.")
    if size == 0:
        raise _empty()
    with open(path, "rb") as fh:
        head = fh.read(8)
    if suffix in TEXT_SUFFIXES:
        if head.startswith(_ZIP_MAGIC) or head.startswith(_CFB_MAGIC):
            raise UserError("Filen er et Excel-regneark med feil filendelse. Gi den endelsen .xlsx eller .xls.")
        wb = _read_csv(path)
    elif head.startswith(_CFB_MAGIC):
        # Encrypted .xlsx files are OLE containers with an EncryptionInfo stream.
        if suffix in EXCEL_SUFFIXES and _ENCRYPTED_MARK in path.read_bytes():
            raise _protected()
        wb = _read_xls(path)
    elif head.startswith(_ZIP_MAGIC):
        wb = _read_xlsx(path)
    else:
        raise _broken()
    if not any(s.height for s in wb.sheets):
        raise _empty()
    return wb


def format_datetime(v: datetime) -> str:
    if v.hour == v.minute == v.second == 0 and not v.microsecond:
        return v.strftime("%Y-%m-%d")
    if v.second == 0:
        return v.strftime("%Y-%m-%d %H:%M")
    return v.strftime("%Y-%m-%d %H:%M:%S")

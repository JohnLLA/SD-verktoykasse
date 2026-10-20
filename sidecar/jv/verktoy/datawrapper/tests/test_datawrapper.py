"""The command line against the synthetic files in testdata/datawrapper.

Every fasit is the expected TSV, byte for byte (scripts/lag_testdata_datawrapper.py).
These tests also run against the packed binary (sidecar/build/pytest_packaged.py).
"""

import hashlib
import json
import shutil
from pathlib import Path

import pytest

from jv_datawrapper.cli import main

DATA = Path(__file__).resolve().parents[3] / "testdata" / "datawrapper"


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def run(capsys, *args):
    code = main([str(a) for a in args])
    captured = capsys.readouterr()
    report = json.loads(captured.out) if code == 0 else None
    return code, report, captured


def fasit(name):
    return (DATA / name).read_bytes().decode("utf-8")


CASES = [
    ("ssb_befolkning.xlsx", [], "ssb_befolkning.fasit.tsv"),
    ("ssb_befolkning.xlsx", ["--med-sum"], "ssb_befolkning.med-sum.fasit.tsv"),
    ("to_overskriftsrader.xlsx", [], "to_overskriftsrader.fasit.tsv"),
    ("sumrader.xlsx", [], "sumrader.fasit.tsv"),
    ("sumrader.xlsx", ["--med-sum"], "sumrader.med-sum.fasit.tsv"),
    ("to_tabeller.xlsx", [], "to_tabeller.fasit.tsv"),
    ("to_tabeller.xlsx", ["--tabell", "1"], "to_tabeller.tabell1.fasit.tsv"),
    ("to_tabeller.xlsx", ["--ark", "Ark1", "--tabell", "2"], "to_tabeller.fasit.tsv"),
    ("skjult_rad.xlsx", [], "skjult_rad.fasit.tsv"),
    ("tall_som_tekst.xlsx", [], "tall_som_tekst.fasit.tsv"),
    ("semikolon_cp1252.csv", [], "semikolon_cp1252.fasit.tsv"),
    ("bytt.csv", [], "bytt.fasit.tsv"),
    ("bytt.csv", ["--bytt"], "bytt.bytt.fasit.tsv"),
    ("omsetning.xls", [], "omsetning.fasit.tsv"),
]


@pytest.mark.parametrize("name, extra, expected", CASES, ids=[f"{c[0]}{''.join(c[1])}" for c in CASES])
def test_tsv_matches_fasit_byte_for_byte(capsys, name, extra, expected):
    before = sha(DATA / name)
    code, report, captured = run(capsys, "lag", DATA / name, *extra, "--json")
    assert code == 0, captured.err
    assert report["tool"] == "datawrapper"
    assert report["datawrapper"]["tsv"].encode("utf-8") == fasit(expected).encode("utf-8")
    assert report["outputs"] == []
    assert sha(DATA / name) == before  # the original is never changed


def test_ssb_title_source_notes_symbols_and_sum(capsys):
    code, report, _ = run(capsys, "lag", DATA / "ssb_befolkning.xlsx")
    assert code == 0
    d = report["datawrapper"]
    assert d["title"] == "Tabell 1. Folkemengde etter kommune. 1. januar"
    assert d["description"] == "Fiktive tall til test"
    assert d["source"] == "Statistisk sentralbyrå (fiktive tall)"
    assert d["notes"][:2] == ["¹ Ny kommune fra 2024.", "* Foreløpige tall."]
    assert d["notes"][2].startswith("Tegnforklaring")
    assert d["symbols"] == {".": 1, "..": 1, ":": 1, "-": 1, "x": 1}
    assert d["removed"]["sum_rows"] == ["Hele landet"]
    assert d["header_rows"] == 1
    assert [c["name"] for c in d["columns"]][:2] == ["Kommune", "2022"]
    assert [c["numeric"] for c in d["columns"]] == [False, True, True, True, True]
    assert d["rows"] == 5
    assert any("sumrad" in w for w in report["warnings"])
    assert any("«..»" in w for w in report["warnings"])
    assert report["summary"].startswith("Tabellen har 5 rader og 5 kolonner")


def test_two_header_rows_are_joined(capsys):
    code, report, _ = run(capsys, "lag", DATA / "to_overskriftsrader.xlsx")
    d = report["datawrapper"]
    assert d["header_rows"] == 2
    assert d["title"] == "Sysselsatte etter fylke og kjønn"
    assert d["source"] == "SSB (fiktive tall)"


def test_hierarchy_and_sums_are_reported(capsys):
    code, report, _ = run(capsys, "lag", DATA / "sumrader.xlsx")
    d = report["datawrapper"]
    assert d["filled_labels"] == 4
    assert d["removed"]["sum_rows"] == ["Nordfylke – Sum", "Sørfylke – Sum", "I alt"]
    assert d["removed"]["sum_columns"] == ["Totalt"]
    code, report, _ = run(capsys, "lag", DATA / "sumrader.xlsx", "--med-sum")
    d = report["datawrapper"]
    assert d["removed"]["sum_rows"] == [] and d["removed"]["sum_rows_found"]
    assert [c["sum"] for c in d["columns"]] == [False, False, False, False, True]


def test_hidden_rows_columns_and_sheets(capsys):
    code, report, _ = run(capsys, "lag", DATA / "skjult_rad.xlsx")
    d = report["datawrapper"]
    assert d["sheet"] == "Data"
    assert d["hidden_sheets"] == 1
    assert d["removed"]["hidden_rows"] == 1
    assert d["removed"]["hidden_columns"] == 1
    code, report, _ = run(capsys, "lag", DATA / "skjult_rad.xlsx", "--ark", "Oversikt")
    assert report["datawrapper"]["tsv"] == "Nøkkeltall\tVerdi\nKommuner\t3\nInnbyggere\t3600\n"
    code, _, captured = run(capsys, "lag", DATA / "skjult_rad.xlsx", "--ark", "Hjelpeark")
    assert code == 1 and "Fant ikke arket" in captured.err


def test_numbers_as_text_warnings(capsys):
    code, report, _ = run(capsys, "lag", DATA / "tall_som_tekst.xlsx")
    d = report["datawrapper"]
    assert d["uncertain_cells"] == ["C4", "D4"]
    text = " ".join(report["warnings"])
    assert "to måter" in text and "C4, D4" in text
    assert "prosenttegnet i 2 celler" in text
    assert "formatert som prosent" not in text  # Excel percent is converted silently (12.5)
    assert [c["name"] for c in d["columns"]][4] == "Andel (%)"
    assert "feilverdi fra Excel" in text


def test_csv_detection_is_reported(capsys):
    code, report, _ = run(capsys, "analyser", DATA / "semikolon_cp1252.csv")
    d = report["datawrapper"]
    assert d["format"] == "csv"
    assert d["encoding"] == "Windows-1252"
    assert d["delimiter"] == "semikolon"
    code, report, _ = run(capsys, "analyser", DATA / "bytt.csv")
    assert report["datawrapper"]["encoding"] == "UTF-8 med BOM"
    assert report["datawrapper"]["delimiter"] == "komma"


def test_analyser_lists_sheets_tables_and_suggestions(capsys):
    code, report, _ = run(capsys, "analyser", DATA / "to_tabeller.xlsx", "--json")
    assert code == 0
    d = report["datawrapper"]
    assert d["suggested_sheet"] == "Ark1"
    (sheet,) = d["sheets"]
    assert sheet["suggested_table"] == 2
    t1, t2 = sheet["tables"]
    assert (t1["index"], t1["range"], t1["suggested"]) == (1, "A2:C5", False)
    assert (t2["index"], t2["range"], t2["suggested"]) == (2, "A9:C13", True)
    assert t1["title"] == "Tabell 1. Besøkende per måned"
    assert t2["headers"] == ["Type", "Antall", "Andel"]
    assert t2["source"] == "Fiktive tall"
    assert report["summary"] == "Fant 2 tabeller i 1 ark."


def test_analyser_picks_sheet_with_most_numbers_and_skips_hidden(capsys):
    code, report, _ = run(capsys, "analyser", DATA / "skjult_rad.xlsx")
    d = report["datawrapper"]
    assert [s["name"] for s in d["sheets"]] == ["Oversikt", "Data"]
    assert d["suggested_sheet"] == "Data"
    assert d["hidden_sheets"] == 1
    assert any("skjult ark" in w for w in report["warnings"])


def test_transpose_is_suggested(capsys):
    code, report, _ = run(capsys, "analyser", DATA / "bytt.csv")
    assert report["datawrapper"]["sheets"][0]["tables"][0]["transpose_suggested"] is True
    code, report, _ = run(capsys, "lag", DATA / "bytt.csv")
    assert report["datawrapper"]["transpose_suggested"] is True
    assert any("bytte rader og kolonner" in w for w in report["warnings"])
    code, report, _ = run(capsys, "analyser", DATA / "ssb_befolkning.xlsx")
    assert report["datawrapper"]["sheets"][0]["tables"][0]["transpose_suggested"] is False


def test_remove_columns(capsys):
    code, report, _ = run(capsys, "lag", DATA / "sumrader.xlsx", "--fjern-kolonne", "2", "--fjern-kolonne", "4")
    d = report["datawrapper"]
    assert d["tsv"].splitlines()[:2] == ["Fylke\tBarnehager", "Nordfylke\t12"]
    assert d["removed"]["columns"] == ["Kommune", "Skoler"]
    assert [c["removed"] for c in d["columns"]] == [False, True, False, True]
    code, _, captured = run(capsys, "lag", DATA / "sumrader.xlsx", "--fjern-kolonne", "9")
    assert code == 1 and "Kolonne 9 finnes ikke" in captured.err
    code, _, captured = run(capsys, "lag", DATA / "bytt.csv", "--fjern-kolonne", "1", "2", "3")
    assert code == 1 and "Minst én kolonne" in captured.err


def test_wrong_table_number(capsys):
    code, _, captured = run(capsys, "lag", DATA / "to_tabeller.xlsx", "--tabell", "3")
    assert code == 1
    assert "har 2 tabeller" in captured.err


def test_csv_file_utf8_comma_and_original_untouched(capsys, tmp_path):
    target = tmp_path / "ut.csv"
    code, report, _ = run(capsys, "lag", DATA / "ssb_befolkning.xlsx", "--csv", target)
    assert code == 0
    assert report["outputs"] == [{"path": str(target), "type": "csv"}]
    raw = target.read_bytes()
    assert not raw.startswith(b"\xef\xbb\xbf")
    lines = raw.decode("utf-8").splitlines()
    assert lines[0] == "Kommune,2022,2023,2024,Endring 2023–2024 (%)"
    assert lines[1] == "3101 Nordby,12345,12401,12502,0.8"
    assert lines[2] == "3102 Sørvik¹,8210,8199,,"
    # An existing file is never overwritten.
    code, report, _ = run(capsys, "lag", DATA / "sumrader.xlsx", "--csv", target)
    assert report["outputs"][0]["path"] == str(tmp_path / "ut-2.csv")
    assert target.read_bytes() == raw


def test_csv_next_to_a_csv_original_never_replaces_it(capsys, tmp_path):
    src = tmp_path / "data.csv"
    shutil.copy(DATA / "semikolon_cp1252.csv", src)
    before = sha(src)
    code, report, _ = run(capsys, "lag", src, "--csv", src)
    assert code == 0
    assert report["outputs"][0]["path"] == str(tmp_path / "data-2.csv")
    assert sha(src) == before
    assert (tmp_path / "data-2.csv").read_text(encoding="utf-8").startswith("Kommune,Innbyggere,Andel over 67 år,Endring\n")


def test_csv_must_end_in_csv_and_folder_must_exist(capsys, tmp_path):
    code, _, captured = run(capsys, "lag", DATA / "bytt.csv", "--csv", tmp_path / "ut.txt")
    assert code == 1 and ".csv" in captured.err
    code, _, captured = run(capsys, "lag", DATA / "bytt.csv", "--csv", tmp_path / "finnes-ikke" / "ut.csv")
    assert code == 1 and "finnes ikke" in captured.err


@pytest.mark.parametrize(
    "name, message",
    [
        ("tom.xlsx", "Filen er tom"),
        ("tom.csv", "Filen er tom"),
        ("odelagt.xlsx", "ødelagt"),
        ("passord.xlsx", "passordbeskyttet"),
    ],
)
def test_empty_broken_and_protected_files(capsys, name, message):
    for cmd in ("analyser", "lag"):
        code, report, captured = run(capsys, cmd, DATA / name)
        assert code == 1
        assert message in captured.err
        assert "Traceback" not in captured.err


def test_too_large_and_wrong_type(capsys, tmp_path):
    big = tmp_path / "stor.csv"
    with open(big, "wb") as fh:
        fh.truncate(50 * 1024 * 1024 + 1)
    code, _, captured = run(capsys, "analyser", big)
    assert code == 1 and "50 MB" in captured.err
    ods = tmp_path / "regneark.ods"
    ods.write_bytes(b"PK\x03\x04")
    code, _, captured = run(capsys, "analyser", ods)
    assert code == 1 and "støttes ikke" in captured.err
    code, _, captured = run(capsys, "analyser", tmp_path / "finnes-ikke.xlsx")
    assert code == 1 and "Finner ikke" in captured.err


def test_stderr_has_only_progress_no_cell_content(capsys):
    code, report, captured = run(capsys, "lag", DATA / "ssb_befolkning.xlsx")
    assert code == 0
    lines = [ln for ln in captured.err.splitlines() if ln.strip()]
    assert lines and all(json.loads(ln).keys() == {"progress"} for ln in lines)
    assert json.loads(lines[-1])["progress"] == 1.0
    for content in ("Nordby", "Folkemengde", "Statistisk"):
        assert content not in captured.err


def test_bad_arguments_are_user_errors(capsys):
    code, _, captured = run(capsys, "lag", DATA / "bytt.csv", "--tabell", "en")
    assert code == 1 and "heltall" in captured.err
    code, _, captured = run(capsys, "finn", DATA / "bytt.csv")
    assert code == 1 and "Ugyldige argumenter" in captured.err


# ------------------------------------------------------------------ hostile files
# Built in the test (tmp_path), never stored in git.


def _xlsx(path, rows=(("Navn", "Tall"), ("a", 1), ("b", 2))):
    from openpyxl import Workbook

    wb = Workbook()
    for row in rows:
        wb.active.append(list(row))
    wb.save(path)
    return path


def _rewrite(src, dst, changes):
    """Copy a zip, replacing (or adding) members: {name: bytes or function(old bytes)}."""
    import zipfile

    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        names = [i.filename for i in zin.infolist()]
        for name in names:
            data = zin.read(name)
            change = changes.get(name)
            if callable(change):
                data = change(data)
            elif change is not None:
                data = change
            zout.writestr(name, data)
        for name, change in changes.items():
            if name not in names:
                zout.writestr(name, change)
    return dst


def test_one_cell_far_out_is_refused_quickly(capsys, tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Navn", "Tall"])
    ws.append(["a", 1])
    ws["XFD1048576"] = "langt ute"
    f = tmp_path / "glissen.xlsx"
    wb.save(f)
    for cmd in ("analyser", "lag"):
        code, _, captured = run(capsys, cmd, f)
        assert code == 1
        assert "for stort" in captured.err


def test_zip_bomb_is_refused_before_unpacking(capsys, tmp_path):
    good = _xlsx(tmp_path / "god.xlsx")
    # 64 MB of zeros packs to about 64 kB: a ratio far above 100:1.
    bomb = _rewrite(good, tmp_path / "bombe.xlsx", {"xl/media/fyll.bin": b"\0" * (64 * 1024 * 1024)})
    assert bomb.stat().st_size < 1024 * 1024
    code, _, captured = run(capsys, "analyser", bomb)
    assert code == 1
    assert "pakker seg ut" in captured.err


BILLION_LAUGHS = (
    b'<?xml version="1.0"?>\n<!DOCTYPE lolz [<!ENTITY lol "lol">'
    b'<!ENTITY lol2 "&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;&lol;">'
    b'<!ENTITY lol3 "&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;&lol2;">'
    b'<!ENTITY lol4 "&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;&lol3;">]>\n'
)


def _body(data):
    return data.split(b"?>", 1)[1] if data.startswith(b"<?xml") else data


def test_xml_entities_are_not_expanded(capsys, tmp_path):
    """No part of an .xlsx may carry a DTD: openpyxl reads sheets with defusedxml,
    but other parts with lxml, which expands internal entities."""
    good = _xlsx(tmp_path / "god.xlsx", rows=(("Navn", "Tall"), ("ENTITET", 1), ("b", 2)))

    def in_sheet(data):
        return BILLION_LAUGHS + _body(data).replace(b"ENTITET", b"&lol4;")

    f = _rewrite(good, tmp_path / "entitet-ark.xlsx", {"xl/worksheets/sheet1.xml": in_sheet})
    code, _, captured = run(capsys, "lag", f)
    assert code == 1
    assert "lollol" not in captured.out + captured.err

    def in_workbook(data):
        return BILLION_LAUGHS + _body(data).replace(b'name="Sheet"', b'name="&lol4;"')

    f = _rewrite(good, tmp_path / "entitet-bok.xlsx", {"xl/workbook.xml": in_workbook})
    code, _, captured = run(capsys, "analyser", f)
    assert code == 1 and "DTD" in captured.err
    assert "lollol" not in captured.out + captured.err

    def comment_first(data):  # a long comment must not hide a DTD after it
        return b"<!--" + b"x" * 70000 + b"-->" + BILLION_LAUGHS + _body(data)

    f = _rewrite(good, tmp_path / "entitet-kommentar.xlsx", {"xl/workbook.xml": comment_first})
    code, _, captured = run(capsys, "analyser", f)
    assert code == 1


def test_external_entities_are_not_read(capsys, tmp_path):
    secret = tmp_path / "hemmelig.txt"
    secret.write_text("HEMMELIG-INNHOLD", encoding="utf-8")
    good = _xlsx(tmp_path / "god.xlsx", rows=(("Navn", "Tall"), ("ENTITET", 1), ("b", 2)))

    def xxe(data):
        doctype = f'<?xml version="1.0"?>\n<!DOCTYPE worksheet [<!ENTITY xxe SYSTEM "{secret.as_uri()}">]>\n'
        return doctype.encode() + _body(data).replace(b"ENTITET", b"&xxe;")

    f = _rewrite(good, tmp_path / "xxe.xlsx", {"xl/worksheets/sheet1.xml": xxe})
    code, _, captured = run(capsys, "lag", f)
    assert code == 1
    assert "HEMMELIG" not in captured.out + captured.err


def test_messages_do_not_name_files_sheets_or_folders(capsys, tmp_path):
    code, _, captured = run(capsys, "lag", DATA / "skjult_rad.xlsx", "--ark", "Hjelpeark")
    assert code == 1 and "Hjelpeark" not in captured.err
    code, _, captured = run(capsys, "lag", DATA / "bytt.csv", "--csv", tmp_path / "Hemmelig mappe" / "ut.csv")
    assert code == 1 and "Hemmelig" not in captured.err
    code, _, captured = run(capsys, "analyser", tmp_path / "Kildens navn.xlsx")
    assert code == 1 and "Kildens" not in captured.err
    code, _, captured = run(capsys, "analyser", tmp_path / "Kildens navn.ods")
    assert code == 1 and "Kildens" not in captured.err


LONG = b"x" * 70000


@pytest.mark.parametrize(
    "prefix",
    [
        b"<!--<a-->" + b"<!--" + LONG + b"-->",  # «<a» inside a comment is not the root
        b"<?pi <a ?>" + b"<!--" + LONG + b"-->",  # nor inside a processing instruction
        b"<!--<a " + LONG + b"-->",
    ],
    ids=["kommentar", "pi", "lang-kommentar"],
)
def test_dtd_hidden_after_fake_roots_is_refused(capsys, tmp_path, prefix):
    good = _xlsx(tmp_path / "god.xlsx")

    def core(data):
        return b'<?xml version="1.0"?>' + prefix + BILLION_LAUGHS.split(b"?>", 1)[1] + _body(data)

    f = _rewrite(good, tmp_path / "skjult-dtd.xlsx", {"docProps/core.xml": core})
    code, _, captured = run(capsys, "analyser", f)
    assert code == 1 and "DTD" in captured.err


def test_part_without_root_within_the_limit_is_refused(capsys, tmp_path):
    good = _xlsx(tmp_path / "god.xlsx")
    f = _rewrite(
        good,
        tmp_path / "uten-rot.xlsx",
        {"docProps/app.xml": b'<?xml version="1.0"?>' + b"<!--" + b"x" * (300 * 1024) + b"--><a/>"},
    )
    code, _, captured = run(capsys, "analyser", f)
    assert code == 1 and "DTD" in captured.err


def test_utf16_dtd_is_refused(capsys, tmp_path):
    good = _xlsx(tmp_path / "god.xlsx")
    xml = '<?xml version="1.0" encoding="UTF-16"?><!DOCTYPE a [<!ENTITY e "x">]><a>&e;</a>'
    for name, data in (
        ("med-bom", "﻿".encode("utf-16-le") + xml.encode("utf-16-le")),
        ("uten-bom", xml.encode("utf-16-be")),
    ):
        f = _rewrite(good, tmp_path / f"utf16-{name}.xlsx", {"docProps/custom.xml": data})
        code, _, captured = run(capsys, "analyser", f)
        assert code == 1 and "DTD" in captured.err, name


def test_packed_binary_parses_xml_with_defusedxml_without_lxml(capsys):
    code, report, _ = run(capsys, "analyser", DATA / "sumrader.xlsx")
    assert code == 0
    assert report["datawrapper"]["xml_parser"] == {"lxml": False, "defusedxml": True}

"""Core logic: translate .txt, .md and .docx between bokmål and nynorsk with Apertium.

Apertium is called through subprocess with argument lists, in the standard
mode that marks problem words in its output: ``*word`` (unknown word),
``#word`` or ``in#word`` (lemma known, but the form cannot be generated) and
``@word`` (missing in the bilingual dictionary). Real Apertium moves and
duplicates ``*``, ``#`` and ``@`` that stand in the author's text (Markdown
``**bold**``, hashtags, ``@names``), also when they are disguised, so those
characters never reach Apertium: the text is cut at every run of them, the
pieces are translated, and the author's characters are put back between the
translated pieces. Every ``*``, ``#`` and ``@`` in Apertium's output is
therefore a mark. Each marked word is wrapped as ``[?word]`` and listed; no raw
mark is ever left in the draft.

No printing, no sys.exit: errors are raised as UserError.
"""

from __future__ import annotations

import os
import re
import shlex
import subprocess
import sys
import tempfile
import threading
import zipfile
from collections.abc import Callable, Iterable
from dataclasses import dataclass, field
from pathlib import Path

from jv_common import UserError, safe_output_path

# --to value -> Apertium mode
MODES = {"nn": "nob-nno", "nb": "nno-nob"}
TARGET_NAMES = {"nn": "nynorsk", "nb": "bokmål"}
SUPPORTED_SUFFIXES = (".txt", ".md", ".docx")

MARK_TYPES = {"*": "unknown", "#": "generation", "@": "untranslated"}
MARK_TYPES_NO = {
    "unknown": "ukjent ord",
    "generation": "bøyingsfeil",
    "untranslated": "uten oversettelse",
}
MARKER_TEMPLATE = "[?{word}]"

SETUP_DOC = "verktoy/nynorsk/OPPSETT.md"
TIMEOUT_SECONDS = 900

_WORD = re.compile(r"\w+(?:[-'’]\w+)*")

ProgressFn = Callable[[float], None]


# --------------------------------------------------------------------------
# Apertium: the mode pipeline, run directly (no shell, no apertium script)
# --------------------------------------------------------------------------
#
# ``apertium -f txt <mode>`` is a bash script that runs
#
#     apertium-destxt | <mode file, rewritten by apertium-wblank-mode> | apertium-retxt
#
# with ``$1`` = ``-g`` (keep the *, # and @ marks) and ``$2`` empty, plus two
# gawk filters for AP_SETVAR, which is empty here and then passes the text on
# unchanged. The tool does the same itself: it reads the mode file, starts the
# programs with subprocess and connects them with pipes. That needs neither
# bash nor gawk, so it works the same on macOS and Windows, and the output is
# byte for byte what the script gives.
#
# Safety: a stage may only start one of the Apertium programs below, and a
# file argument must be a relative path inside the data folder. The programs
# run with the data folder as working directory and get only those relative
# ASCII paths, so a user name with æøå in the path cannot break them (on
# Windows they open files with the ANSI code page). The text goes on stdin.

PIPELINE_PROGRAMS = frozenset({
    "apertium-destxt", "apertium-retxt", "apertium-wblank-attach", "apertium-wblank-detach",
    "lt-proc", "lt-merge", "cg-proc", "apertium-tagger", "lsx-proc", "apertium-pretransfer",
    "lrx-proc", "apertium-transfer", "apertium-interchunk", "apertium-postchunk",
})
_OPTION = re.compile(r"--?[A-Za-z0-9][A-Za-z0-9-]*")
_REL_DATA = re.compile(r"[A-Za-z0-9._+-]+(?:/[A-Za-z0-9._+-]+)*")
# Values of the mode's positional parameters for ``apertium -f txt``.
_PARAMS = {"$1": "-g", "$2": None}


def _missing_apertium_message() -> str:
    return (
        "Apertium er ikke installert eller ikke i PATH. Verktøyet trenger "
        "Apertium-programmene (lttoolbox, apertium, apertium-lex-tools, "
        "apertium-separable og cg3) og språkparet apertium-nno-nob (modusene "
        "nob-nno og nno-nob). Installer begge og prøv igjen; se "
        f"{SETUP_DOC} for fremgangsmåte på macOS (bygging fra kilde)."
    )


def _missing_mode_message(mode: str) -> str:
    return (
        f"Apertium er installert, men språkretningen «{mode}» finnes ikke. "
        "Installer språkparet apertium-nno-nob (gir både nob-nno og nno-nob); "
        f"se {SETUP_DOC}."
    )


class PipelineError(RuntimeError):
    """The installation is not what the tool expects (internal error, code 2)."""


def _program_argv(name: str) -> list[str] | None:
    """Command for an Apertium program, found in PATH (never the working directory).

    The packages put their own program folder first in PATH (jv-package.json).
    """
    exts = [""]
    if os.name == "nt":
        exts = [e.lower() for e in os.environ.get("PATHEXT", ".EXE").split(os.pathsep) if e]
    for folder in os.environ.get("PATH", "").split(os.pathsep):
        if not folder or not os.path.isabs(folder):
            continue
        for ext in exts:
            candidate = os.path.join(folder, name + ext)
            if os.path.isfile(candidate) and os.access(candidate, os.X_OK):
                return [candidate]
    return None


def find_datadir() -> Path | None:
    """Apertium's data folder (with modes/): APERTIUM_DATADIR, else next to the programs.

    The packages set APERTIUM_DATADIR. A local install from source
    (<prefix>/bin and <prefix>/share/apertium, see OPPSETT.md) is found from
    where apertium-destxt is.
    """
    value = os.environ.get("APERTIUM_DATADIR")
    if value:
        path = Path(value)
        return path if path.is_absolute() and (path / "modes").is_dir() else None
    found = _program_argv("apertium-destxt")
    if not found:
        return None
    guess = Path(found[-1]).resolve().parent.parent / "share" / "apertium"
    return guess if (guess / "modes").is_dir() else None


def _wblank(stages: list[str]) -> list[str]:
    """What apertium-wblank-mode does to a mode (Apertium 3.9.12, wblank-mode.cc)."""
    joined = " | ".join(stages)
    if "apertium-wblank-" in joined:
        return stages
    restore_caps = "apertium-restore-caps" in joined
    out: list[str] = []
    for stage in stages:
        if "automorf.bin" in stage or "automorf.hfst" in stage:
            out += [stage, "apertium-wblank-attach"]
        elif not restore_caps and ("autogen.bin" in stage or "autogen.hfst" in stage):
            out += ["apertium-wblank-detach", stage]
        else:
            out.append(stage)
    return out


def read_mode(datadir: Path, mode: str) -> list[list[str]]:
    """The stages of a mode as argument lists, with data paths relative to datadir.

    Raises PipelineError if the mode file asks for anything outside the rules.
    """
    text = (datadir / "modes" / f"{mode}.mode").read_text(encoding="utf-8")
    text = " ".join(text.splitlines())
    stages = [s.strip() for s in text.split("|")]
    if not all(stages):
        raise PipelineError("tomt steg i modusfila")
    root = datadir.resolve()
    out: list[list[str]] = [["apertium-destxt"]]
    for stage in _wblank(stages):
        words = shlex.split(stage, posix=True)
        name, args = words[0], []
        if name not in PIPELINE_PROGRAMS:
            raise PipelineError(f"modusfila bruker et program som ikke er tillatt ({name!r})")
        for word in words[1:]:
            if word in _PARAMS:
                if _PARAMS[word] is not None:
                    args.append(_PARAMS[word])
            elif _OPTION.fullmatch(word):
                args.append(word)
            else:
                args.append(_data_path(word, root))
        out.append([name, *args])
    out.append(["apertium-retxt"])
    return out


def _data_path(word: str, root: Path) -> str:
    path = Path(word)
    if path.is_absolute():
        try:
            rel = path.resolve().relative_to(root).as_posix()
        except ValueError:
            raise PipelineError("modusfila peker utenfor datamappa") from None
    else:
        rel = word
    if not _REL_DATA.fullmatch(rel) or ".." in rel.split("/"):
        raise PipelineError("uventet filsti i modusfila")
    if not (root / rel).is_file():
        raise FileNotFoundError(rel)
    return rel


def _child_env() -> dict[str, str]:
    env = dict(os.environ)
    if os.name != "nt":
        # The apertium script always sets a UTF-8 LC_CTYPE (locale_utf8).
        env["LC_CTYPE"] = "C.UTF-8" if sys.platform.startswith("linux") else "en_US.UTF-8"
    return env


class Apertium:
    """One Apertium mode (nob-nno or nno-nob), run as a pipeline of programs."""

    def __init__(self, mode: str, datadir: str | os.PathLike[str] | None = None) -> None:
        found = Path(datadir) if datadir is not None else find_datadir()
        if found is None or _program_argv("lt-proc") is None:
            raise UserError(_missing_apertium_message())
        self.datadir = found
        self.mode = mode
        self._stages: list[list[str]] | None = None
        self._commands: list[list[str]] | None = None

    def available_modes(self) -> list[str]:
        """Modes in <datadir>/modes (what ``apertium -l`` lists)."""
        folder = self.datadir / "modes"
        return sorted(p.stem for p in folder.glob("*.mode")) if folder.is_dir() else []

    def check_mode(self) -> None:
        """Fail early: the mode, its data files and its programs must be there."""
        if self.mode not in self.available_modes():
            raise UserError(_missing_mode_message(self.mode))
        try:
            stages = read_mode(self.datadir, self.mode)
        except FileNotFoundError:
            raise UserError(_missing_mode_message(self.mode)) from None
        commands = []
        for name, *args in stages:
            argv = _program_argv(name)
            if argv is None:
                raise UserError(_missing_apertium_message())
            commands.append([*argv, *args])
        self._stages, self._commands = stages, commands

    def run(self, text: str) -> str:
        """Translate text; problem words come back with *, # and @ marks."""
        if self._commands is None:
            self.check_mode()
        data = text.encode("utf-8")
        out, failures = _run_pipeline(self._commands, data, self.datadir)
        if failures:
            code, stderr = failures[-1]
            detail = stderr.strip().splitlines()[-1] if stderr.strip() else "ingen feilmelding"
            raise UserError(f"Apertium feilet med kode {code}: {detail}")
        if os.name == "nt" and b"\r" not in data:
            # MinGW programs may write \r\n for \n on Windows (stdout in text mode).
            out = out.replace(b"\r\n", b"\n")
        return out.decode("utf-8", errors="replace")


def _run_pipeline(commands: list[list[str]], data: bytes, cwd: Path) -> tuple[bytes, list[tuple[int, str]]]:
    """Run the commands connected by pipes. Returns (stdout, [(code, stderr)] of failed stages).

    stderr of each stage goes to its own temporary file, so no pipe can fill
    up and block. On timeout every stage is stopped.
    """
    flags = getattr(subprocess, "CREATE_NO_WINDOW", 0) if os.name == "nt" else 0
    env = _child_env()
    procs: list[subprocess.Popen] = []
    errs = []
    timed_out = threading.Event()

    def stop_all() -> None:
        timed_out.set()
        for p in procs:
            if p.poll() is None:
                p.kill()

    try:
        for i, argv in enumerate(commands):
            err = tempfile.TemporaryFile()
            errs.append(err)
            try:
                proc = subprocess.Popen(
                    argv,
                    stdin=subprocess.PIPE if i == 0 else procs[-1].stdout,
                    stdout=subprocess.PIPE,
                    stderr=err,
                    cwd=str(cwd),
                    env=env,
                    creationflags=flags,
                )
            except OSError as exc:
                raise UserError(f"Klarte ikke å starte Apertium: {exc.strerror or exc}") from None
            if i > 0:
                procs[-1].stdout.close()  # only the next stage holds the read end
            procs.append(proc)

        def feed() -> None:
            try:
                procs[0].stdin.write(data)
            except (BrokenPipeError, OSError):
                pass
            finally:
                try:
                    procs[0].stdin.close()
                except OSError:
                    pass

        timer = threading.Timer(TIMEOUT_SECONDS, stop_all)
        timer.daemon = True
        timer.start()
        writer = threading.Thread(target=feed, daemon=True)
        writer.start()
        out = procs[-1].stdout.read()
        procs[-1].stdout.close()
        writer.join()
        codes = [p.wait() for p in procs]
        timer.cancel()
        if timed_out.is_set():
            raise UserError(f"Apertium brukte mer enn {TIMEOUT_SECONDS} sekunder og ble stoppet.")
        failures = []
        for code, err in zip(codes, errs):
            if code != 0:
                err.seek(0)
                failures.append((code, err.read().decode("utf-8", errors="replace")))
        return out, failures
    finally:
        for p in procs:
            if p.poll() is None:
                p.kill()
                p.wait()
        for err in errs:
            err.close()


# --------------------------------------------------------------------------
# Marks: Apertium's own *, # and @
# --------------------------------------------------------------------------

MARK_CHARS = "*#@"
# A word with marks: a leading mark, a mark inside (på#drage) or a trailing one.
_MARKED = re.compile(
    r"[*#@]*\w+(?:[-'’#*@]+\w+)*[*#@]*|[*#@]+"
)
# Apertium handles roughly this many tokens before cg-proc forces a break
BATCH_TOKENS = 300
_TOKEN = re.compile(r"\w+|[^\w\s]")


@dataclass
class Translated:
    text: str  # draft with [?word] markers
    marks: list[tuple[str, str]] = field(default_factory=list)  # (word, type)
    offsets: list[int] = field(default_factory=list)  # marker start in text
    warnings: list[str] = field(default_factory=list)


def parse_marks(output: str) -> Translated:
    """Turn Apertium output with raw marks into a draft with [?word] markers.

    Marks from the author must already be protected (see translate_text), so
    every ``*``, ``#`` and ``@`` here comes from Apertium. When a word has
    several marks the first one decides the type.
    """
    pieces: list[str] = []
    marks: list[tuple[str, str]] = []
    offsets: list[int] = []
    length = 0
    cursor = 0
    for match in _MARKED.finditer(output):
        chunk = match.group(0)
        found = [c for c in chunk if c in MARK_CHARS]
        if not found:
            continue
        before = output[cursor:match.start()]
        word = "".join(c for c in chunk if c not in MARK_CHARS)
        cursor = match.end()
        if not word:
            pieces.append(before)  # a lone mark with no word: drop it
            length += len(before)
            continue
        marker = MARKER_TEMPLATE.format(word=word)
        pieces += [before, marker]
        length += len(before)
        offsets.append(length)
        length += len(marker)
        marks.append((word, MARK_TYPES[found[0]]))
    pieces.append(output[cursor:])
    return Translated("".join(pieces), marks, offsets)


_SYMBOL_RUNS = re.compile(r"([*#@]+)")
_PADDING = re.compile(r"(\s*)(.*?)(\s*)\Z", re.DOTALL)


def _translate_cores(engine: Apertium, cores: list[str]) -> tuple[list[str], int]:
    """Translate pieces of text, one output per input (joined by blank lines).

    Returns (raw outputs, number of pieces left untranslated). A piece whose
    line count changes is retried alone, then line by line; if a line still changes, the original
    line is kept so nothing is silently damaged.
    """
    counts = [c.count("\n") + 1 for c in cores]
    # A blank line between pieces: with a single newline real Apertium
    # sometimes adds an extra one after words like "var".
    lines = engine.run("\n\n".join(cores)).rstrip("\n").split("\n")
    if len(lines) == sum(counts) + len(cores) - 1:
        out, pos, intact = [], 0, True
        for n in counts:
            out.append("\n".join(lines[pos:pos + n]))
            intact = intact and (pos + n >= len(lines) or lines[pos + n] == "")
            pos += n + 1
        if intact:
            return out, 0
    out, skipped = [], 0
    for core, n in zip(cores, counts):
        single = engine.run(core).rstrip("\n").split("\n")
        if len(single) == n:
            out.append("\n".join(single))
            continue
        done = []  # still wrong: one call per line
        for line in core.split("\n"):
            if not line.strip():
                done.append(line)
                continue
            result = engine.run(line).rstrip("\n").split("\n")
            if len(result) == 1:
                done.append(result[0])
            else:
                done.append(line)
                skipped += 1
        out.append("\n".join(done))
    return out, skipped


def translate_text(engine: Apertium, text: str) -> Translated:
    """Translate text; the author's own *, # and @ stay exactly where they are."""
    parts = _SYMBOL_RUNS.split(text)  # even index: text, odd index: symbols
    padded: dict[int, tuple[str, str, str]] = {}
    for i in range(0, len(parts), 2):
        lead, core, trail = _PADDING.match(parts[i]).groups()
        padded[i] = (lead, core, trail)
    indexes = [i for i, (_l, core, _t) in padded.items() if core]
    outputs: dict[int, str] = {}
    skipped = 0
    if indexes:
        raws, skipped = _translate_cores(engine, [padded[i][1] for i in indexes])
        outputs = dict(zip(indexes, raws))

    pieces: list[str] = []
    marks: list[tuple[str, str]] = []
    offsets: list[int] = []
    length = 0

    def add(piece: str) -> None:
        nonlocal length
        pieces.append(piece)
        length += len(piece)

    for i, part in enumerate(parts):
        if i % 2:
            add(part)
            continue
        lead, core, trail = padded[i]
        add(lead)
        if core:
            res = parse_marks(outputs[i])
            offsets.extend(length + o for o in res.offsets)
            marks.extend(res.marks)
            add(res.text)
        add(trail)
    warnings = []
    if skipped:
        warnings.append(
            f"{skipped} linjer ble ikke oversatt fordi Apertium endret "
            "linjestrukturen. De står uendret i utkastet."
        )
    return Translated("".join(pieces), marks, offsets, warnings)


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
_TEXT_RUNS_XPATH = (
    "./w:r | ./w:hyperlink/w:r | ./w:ins/w:r | ./w:smartTag/w:r "
    "| ./w:fldSimple/w:r"
)


def _is_line_break(el) -> bool:
    return el.tag in (f"{_W}br", f"{_W}cr") and el.get(f"{_W}type") in (
        None,
        "textWrapping",
    )


def _run_text(r) -> str:
    parts: list[str] = []
    for child in r:
        if child.tag == f"{_W}t":
            parts.append(child.text or "")
        elif child.tag == f"{_W}tab":
            parts.append("\t")
        elif _is_line_break(child):
            parts.append("\n")
    return "".join(parts)


def _text_children(r) -> list:
    return [
        c
        for c in r
        if c.tag in (f"{_W}t", f"{_W}tab") or _is_line_break(c)
    ]


class _DocxParagraph:
    """A paragraph whose text lives in a list of w:r elements."""

    def __init__(self, paragraph) -> None:
        self.paragraph = paragraph
        self.runs = [r for r in paragraph._p.xpath(_TEXT_RUNS_XPATH)]
        self.text = "".join(_run_text(r) for r in self.runs)

    def mixed_formatting(self) -> bool:
        from lxml import etree  # dependency of python-docx

        formats = set()
        for r in self.runs:
            if not _run_text(r).strip():
                continue
            rpr = r.find(f"{_W}rPr")
            formats.add(b"" if rpr is None else etree.tostring(rpr))
        return len(formats) > 1

    def set_text(self, text: str) -> None:
        from docx.text.run import Run

        first = next((r for r in self.runs if _run_text(r)), None)
        if first is None:
            return
        for r in self.runs:
            if r is first:
                continue
            for child in _text_children(r):
                r.remove(child)
        # Run.text replaces w:t/w:tab/w:br children and turns \t and \n into
        # w:tab and w:br; keep non-text children such as drawings.
        for child in _text_children(first):
            first.remove(child)
        Run(first, self.paragraph).text = text


def _iter_table_paragraphs(tables, seen: set[int]):
    for table in tables:
        for row in table.rows:
            for cell in row.cells:
                if id(cell._tc) in seen:
                    continue  # merged cells are returned more than once
                seen.add(id(cell._tc))
                yield from cell.paragraphs
                yield from _iter_table_paragraphs(cell.tables, seen)


def _iter_docx_paragraphs(document):
    yield from document.paragraphs
    yield from _iter_table_paragraphs(document.tables, set())


def _open_docx(path: Path):
    import docx

    if not zipfile.is_zipfile(path):
        raise UserError(
            f"«{path.name}» er ikke en gyldig DOCX-fil (ikke et Word-dokument)."
        )
    try:
        return docx.Document(str(path))
    except Exception:  # noqa: BLE001 - python-docx raises many types
        raise UserError(
            f"Klarte ikke å lese «{path.name}» som Word-dokument (DOCX)."
        ) from None


def _batches(paragraphs: list[_DocxParagraph]) -> list[list[_DocxParagraph]]:
    """Group paragraphs so each Apertium call stays below the cohort limit."""
    batches: list[list[_DocxParagraph]] = []
    current: list[_DocxParagraph] = []
    size = 0
    for p in paragraphs:
        tokens = len(_TOKEN.findall(p.text)) + 1
        if current and size + tokens > BATCH_TOKENS:
            batches.append(current)
            current, size = [], 0
        current.append(p)
        size += tokens
    if current:
        batches.append(current)
    return batches


def _translate_batch(
    engine: Apertium, paragraphs: list[_DocxParagraph]
) -> list[Translated] | None:
    """Translate paragraphs in one call (joined by blank lines).

    Returns None if Apertium did not keep the line structure, so the caller
    can fall back to one call per paragraph.
    """
    line_counts = [p.text.count("\n") + 1 for p in paragraphs]
    joined = "\n\n".join(p.text for p in paragraphs)  # blank line = safe separator
    result = translate_text(engine, joined)
    lines = result.text.split("\n")
    if len(lines) != sum(line_counts) + len(paragraphs) - 1:
        return None
    mark_lines = [result.text.count("\n", 0, off) for off in result.offsets]
    out: list[Translated] = []
    first_line = 0
    for count in line_counts:
        last_line = first_line + count
        if last_line < len(lines) and lines[last_line] != "":
            return None
        chunk = "\n".join(lines[first_line:last_line])
        para_marks = [
            m
            for m, line in zip(result.marks, mark_lines)
            if first_line <= line < last_line
        ]
        out.append(Translated(chunk, para_marks))
        first_line = last_line + 1
    if out:
        out[0].warnings = result.warnings
    return out


def _translate_lines(
    engine: Apertium,
    paragraphs: list[_DocxParagraph],
    progress: ProgressFn | None = None,
) -> list[Translated]:
    """Translate in batches; a batch that loses its structure goes paragraph by paragraph."""
    results: list[Translated] = []
    done = 0
    for batch in _batches(paragraphs):
        part = _translate_batch(engine, batch)
        if part is None:
            part = [translate_text(engine, p.text) for p in batch]
        results.extend(part)
        done += len(batch)
        if progress:
            progress(0.9 * done / len(paragraphs))
    return results


def translate_docx(
    engine: Apertium, src: Path, dst: Path, progress: ProgressFn | None = None
) -> tuple[list[tuple[str, str]], list[str], int]:
    """Translate paragraph by paragraph. Returns (marks, warnings, paragraphs)."""
    document = _open_docx(src)
    paragraphs = [
        p
        for p in (_DocxParagraph(par) for par in _iter_docx_paragraphs(document))
        if p.text.strip()
    ]
    if not paragraphs:
        raise UserError(f"«{src.name}» inneholder ingen tekst å oversette.")

    warnings: list[str] = []
    results = _translate_lines(engine, paragraphs, progress)

    mixed = 0
    marks: list[tuple[str, str]] = []
    for p, res in zip(paragraphs, results):
        if p.mixed_formatting():
            mixed += 1
        p.set_text(res.text)
        marks.extend(res.marks)
        warnings.extend(res.warnings)
    if mixed:
        warnings.append(
            f"{mixed} avsnitt hadde blandet tegnformatering (for eksempel fet "
            "eller kursiv i deler av avsnittet). Den er ikke beholdt: hele "
            "avsnittet har fått formateringen til den første teksten i avsnittet."
        )
    document.save(str(dst))
    return marks, warnings, len(paragraphs)


# --------------------------------------------------------------------------
# Plain text and Markdown
# --------------------------------------------------------------------------


def _read_text(path: Path) -> str:
    data = path.read_bytes()
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        raise UserError(
            f"«{path.name}» er ikke UTF-8-tekst. Lagre filen som UTF-8 og prøv igjen."
        ) from None


# --------------------------------------------------------------------------
# Public entry
# --------------------------------------------------------------------------


@dataclass
class FileResult:
    input: Path
    output: Path
    mode: str
    paragraphs: int
    marks: list[tuple[str, str]]
    warnings: list[str]

    def unknown_words(self) -> list[dict]:
        counts: dict[tuple[str, str], int] = {}
        for key in self.marks:
            counts[key] = counts.get(key, 0) + 1
        return [
            {"word": w, "type": t, "count": c} for (w, t), c in counts.items()
        ]


def validate_input(path: Path) -> Path:
    if not path.exists():
        raise UserError(f"Filen «{path}» finnes ikke.")
    if not path.is_file():
        raise UserError(f"«{path}» er ikke en fil.")
    if path.suffix.lower() not in SUPPORTED_SUFFIXES:
        raise UserError(
            f"«{path.name}» har et format som ikke støttes. "
            "Bruk .txt, .md eller .docx."
        )
    if path.stat().st_size == 0:
        raise UserError(f"«{path.name}» er tom.")
    return path


def make_engine(to: str, datadir: str | os.PathLike[str] | None = None) -> Apertium:
    if to not in MODES:
        raise UserError("Retningen må være «nn» (til nynorsk) eller «nb» (til bokmål).")
    engine = Apertium(MODES[to], datadir)
    engine.check_mode()
    return engine


def translate_file(
    engine: Apertium,
    to: str,
    path: str | os.PathLike[str],
    out_dir: str | os.PathLike[str],
    progress: ProgressFn | None = None,
) -> FileResult:
    src = validate_input(Path(path))
    suffix = src.suffix.lower()
    if suffix == ".docx":
        _open_docx(src)  # fail early, before an output path is reserved
        dst = safe_output_path(src, out_dir, f"-{to}.docx")
        marks, warnings, paragraphs = translate_docx(engine, src, dst, progress)
    else:
        text = _read_text(src)
        if not text.strip():
            raise UserError(f"«{src.name}» inneholder ingen tekst å oversette.")
        result = translate_text(engine, text)
        dst = safe_output_path(src, out_dir, f"-{to}{suffix}")
        # newline="": the draft keeps the line endings of the original (LF or
        # CRLF). Text mode on Windows would turn the \r\n already in the text
        # into \r\r\n.
        dst.write_text(result.text, encoding="utf-8", newline="")
        marks, warnings = result.marks, result.warnings
        paragraphs = len([b for b in re.split(r"\n\s*\n", text) if b.strip()])
    return FileResult(src, dst, engine.mode, paragraphs, marks, warnings)


def translate_files(
    to: str,
    paths: Iterable[str | os.PathLike[str]],
    out_dir: str | os.PathLike[str],
    progress: ProgressFn | None = None,
    datadir: str | os.PathLike[str] | None = None,
) -> list[FileResult]:
    paths = [Path(p) for p in paths]
    for p in paths:
        validate_input(p)  # all inputs are checked before any work starts
        if p.suffix.lower() == ".docx":
            _open_docx(p)  # a bad file is reported before Apertium is checked
    engine = make_engine(to, datadir)
    results: list[FileResult] = []
    for n, p in enumerate(paths):
        def part(fraction: float, n: int = n) -> None:
            if progress:
                progress((n + fraction) / len(paths))

        results.append(translate_file(engine, to, p, out_dir, part))
        part(1.0)
    return results

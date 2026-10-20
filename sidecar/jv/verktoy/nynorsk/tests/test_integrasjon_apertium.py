"""Integration tests against a real Apertium with apertium-nno-nob.

The tool runs the mode files itself (core.read_mode), so these tests need the
Apertium programs in PATH and the data folder (APERTIUM_DATADIR, or
<prefix>/share/apertium next to the programs). Skipped when they are missing,
or when the language pair is missing. Against a package (pytest_packaged) the
package's programs and data are used.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import docx
import pytest

from jv_nynorsk import cli, core

DATADIR = core.find_datadir()
pytestmark = pytest.mark.skipif(
    DATADIR is None or core._program_argv("lt-proc") is None,
    reason="Apertium mangler; se verktoy/nynorsk/OPPSETT.md",
)
REFERENCE = Path(__file__).resolve().parents[3] / "testdata" / "nynorsk"


def _modes() -> list[str]:
    return core.Apertium("nob-nno").available_modes() if DATADIR else []


@pytest.fixture
def needs_pair():
    modes = _modes()
    missing = [m for m in ("nob-nno", "nno-nob") if m not in modes]
    if missing:
        pytest.skip(
            "Apertium finnes, men språkparet mangler: " + ", ".join(missing)
            + " (se verktoy/nynorsk/OPPSETT.md)"
        )


def _run(capsys, *argv):
    code = cli.main([str(a) for a in argv])
    captured = capsys.readouterr()
    assert code == 0, captured.err
    return json.loads(captured.out)


def _unknown(report) -> dict[str, str]:
    return {u["word"]: u["type"] for u in report["outputs"][0]["unknown_words"]}


def test_precondition_xqztvik_is_unknown_to_real_apertium(needs_pair):
    for mode in ("nob-nno", "nno-nob"):
        out = core.Apertium(mode).run("xqztvik\n")
        assert out.strip() == "*xqztvik"


def test_real_txt_to_nynorsk(needs_pair, testdata, tmp_path, capsys):
    report = _run(capsys, testdata / "avsnitt_nb.txt", "--to", "nn", "--out", tmp_path)
    text = (tmp_path / "avsnitt_nb-nn.txt").read_text("utf-8")
    assert text.count("\n\n") == 2
    assert "frå" in text
    assert report["outputs"][0]["draft"] is True
    for char in "*#@":
        assert char not in text


def test_real_txt_unknown_word_both_directions(needs_pair, tmp_path, capsys):
    for to, text in (("nn", "Han liker xqztvik og skole.\n"), ("nb", "Han likar xqztvik og skule.\n")):
        src = tmp_path / f"ukjent-{to}.txt"
        src.write_text(text, encoding="utf-8")
        report = _run(capsys, src, "--to", to, "--out", tmp_path / "ut")
        out = (tmp_path / "ut" / f"ukjent-{to}-{to}.txt").read_text("utf-8")
        assert "[?xqztvik]" in out and "*" not in out
        assert _unknown(report) == {"xqztvik": "unknown"}


@pytest.mark.parametrize(
    "to,word,expected",
    [
        ("nn", "svibrent", "svibrent"),
        ("nn", "helsehemma", "helsehemma"),
        ("nb", "svibrent", "svibrent"),
        ("nb", "pådrage", "pådrage"),  # Apertium writes på#drage
    ],
)
def test_real_hash_is_marked_and_listed(needs_pair, tmp_path, capsys, to, word, expected):
    src = tmp_path / "hash.txt"
    src.write_text(f"Han skal {word} nå.\n", encoding="utf-8")
    report = _run(capsys, src, "--to", to, "--out", tmp_path / "ut")
    out = (tmp_path / "ut" / f"hash-{to}.txt").read_text("utf-8")
    assert f"[?{expected}]" in out
    for char in "*#@":
        assert char not in out
    assert _unknown(report) == {expected: "generation"}


def test_real_authors_own_marks_are_kept(needs_pair, tmp_path, capsys):
    src = tmp_path / "md.md"
    src.write_text("# Tittel\n\nHan liker **skole** og @navn #sak.\n", encoding="utf-8")
    report = _run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    out = (tmp_path / "ut" / "md-nn.md").read_text("utf-8")
    assert out.startswith("# Tittel\n\n")
    assert "**skule**" in out and "@namn" in out and "#sak" in out
    assert _unknown(report) == {}


@pytest.mark.parametrize("paragraphs", [2, 3])
def test_real_txt_keeps_paragraph_count_and_order(needs_pair, tmp_path, capsys, paragraphs):
    texts = [f"Avsnitt {i} handler om skole og xqztvik{i}." for i in range(paragraphs)]
    src = tmp_path / "avs.txt"
    src.write_text("\n\n".join(texts) + "\n", encoding="utf-8")
    report = _run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    out = (tmp_path / "ut" / "avs-nn.txt").read_text("utf-8")
    parts = out.rstrip("\n").split("\n\n")
    assert len(parts) == paragraphs and out.endswith("\n")
    for i, part in enumerate(parts):
        assert part.startswith(f"Avsnitt {i} ") and f"[?xqztvik{i}]" in part and "skule" in part
    assert report["outputs"][0]["paragraphs"] == paragraphs


def test_real_docx_multiple_paragraphs_in_order(needs_pair, tmp_path, capsys):
    src = tmp_path / "flere.docx"
    document = docx.Document()
    document.add_heading("Overskrift om skole", level=1)
    for i in range(4):
        document.add_paragraph(f"Avsnitt {i} fra skole med xqztvik{i}.")
    document.save(str(src))
    report = _run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    out = docx.Document(str(tmp_path / "ut" / "flere-nn.docx"))
    assert [p.style.name for p in out.paragraphs] == ["Heading 1"] + ["Normal"] * 4
    assert out.paragraphs[0].text == "Overskrift om skule"
    for i, par in enumerate(out.paragraphs[1:]):
        assert par.text == f"Avsnitt {i} frå skule med [?xqztvik{i}]."
    assert set(_unknown(report)) == {f"xqztvik{i}" for i in range(4)}


def test_real_docx_600_short_paragraphs_without_full_stop(needs_pair, tmp_path, capsys):
    src = tmp_path / "liste.docx"
    document = docx.Document()
    for i in range(600):
        extra = " xqztvik" if i % 150 == 3 else ""
        document.add_paragraph(f"Punkt nummer {i} fra{extra}", style="List Bullet")
    document.save(str(src))
    report = _run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    paras = docx.Document(str(tmp_path / "ut" / "liste-nn.docx")).paragraphs
    assert len(paras) == 600
    for i, par in enumerate(paras):
        assert par.text.startswith(f"Punkt nummer {i} frå"), (i, par.text)
        assert ("[?xqztvik]" in par.text) == (i % 150 == 3)
    assert report["outputs"][0]["unknown_words"][0]["count"] == 4


def test_real_txt_600_lines_keep_line_structure(needs_pair, tmp_path, capsys):
    src = tmp_path / "ordliste.txt"
    src.write_text("".join(f"Fra nummer {i}\n" for i in range(600)), encoding="utf-8")
    _run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    lines = (tmp_path / "ut" / "ordliste-nn.txt").read_text("utf-8").split("\n")
    assert lines[-1] == "" and len(lines) == 601
    assert all(l == f"Frå nummer {i}" for i, l in enumerate(lines[:-1]))


def _symbols(text: str) -> list[str]:
    return [c for c in text if c in "*#@"]


def _shape(text: str) -> list[tuple[str, bool, bool]]:
    """Each run of symbols with: space before it, space after it."""
    return [
        (m.group(0), m.start() == 0 or text[m.start() - 1].isspace(),
         m.end() == len(text) or text[m.end()].isspace())
        for m in re.finditer(r"[*#@]+", text)
    ]


@pytest.mark.parametrize(
    "text",
    [
        "Han sa **17-åringen** var ung.",
        "Han sa #2024 var ung",
        "#ærlig **regjeringen**",
        "**sa** #2024",
        "Hei @ola, se saken *nå* også.",
    ],
)
def test_real_authors_symbols_stay_in_place(needs_pair, text):
    from jv_nynorsk.core import make_engine, translate_text

    out = translate_text(make_engine("nn"), text).text
    # same symbols, same order, no extra spaces around them
    assert _shape(out) == _shape(text)


def test_real_random_sentences_keep_symbols(needs_pair, tmp_path, capsys):
    import random

    rng = random.Random(20240101)
    words = ["sa", "var", "ærlig", "en", "regjeringen", "skole", "17-åringen", "fra", "også", "ung"]
    wrap = [lambda w: f"**{w}**", lambda w: f"*{w}*", lambda w: f"#{w}", lambda w: f"@{w}"]
    sentences = []
    for _ in range(100):
        parts = [rng.choice(words) for _ in range(rng.randint(3, 7))]
        for _ in range(rng.randint(1, 2)):
            i = rng.randrange(len(parts))
            parts[i] = rng.choice(wrap)(parts[i])
        sentences.append(" ".join(parts) + ".")
    src = tmp_path / "tilfeldig.txt"
    src.write_text("\n".join(sentences) + "\n", encoding="utf-8")
    _run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    out = (tmp_path / "ut" / "tilfeldig-nn.txt").read_text("utf-8").split("\n")
    assert len(out) == 101
    for before, after in zip(sentences, out):
        assert _shape(after) == _shape(before), (before, after)


@pytest.mark.parametrize("reference", sorted((REFERENCE / "referanse").glob("*.txt")), ids=lambda p: p.name)
def test_real_output_is_byte_for_byte_the_reference(needs_pair, reference):
    # testdata/nynorsk/referanse/<file>.<mode>.txt is what «apertium -f txt <mode>»
    # (the bash script, Apertium 3.9.12, apertium-nno-nob 1.6.0) gave on macOS for
    # testdata/nynorsk/<file>.txt. The pipeline must give exactly the same bytes,
    # on macOS and on Windows.
    name, mode = reference.stem.rsplit(".", 1)
    engine = core.Apertium(mode)
    out = engine.run((REFERENCE / f"{name}.txt").read_text("utf-8"))
    assert out.encode("utf-8") == reference.read_bytes()

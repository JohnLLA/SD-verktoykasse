"""Tests for jv-nynorsk with a fake apertium (see fake_apertium.py)."""

from __future__ import annotations

import hashlib
import json
import shutil

import docx
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH

from jv_common import cli as jv_cli
from jv_nynorsk import cli, core


def sha(path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def run(capsys, *argv):
    code = cli.main([str(a) for a in argv])
    captured = capsys.readouterr()
    report = json.loads(captured.out) if code == 0 else None
    return code, report, captured.err


def progress_lines(err: str) -> list[float]:
    values = []
    for line in err.splitlines():
        if line.startswith("{"):
            values.append(json.loads(line)["progress"])
    return values


# --------------------------------------------------------------------------
# Marks
# --------------------------------------------------------------------------


def test_parse_marks_all_three_kinds_and_inner_hash():
    res = core.parse_marks("ein *knaskelure og #svibrent, på#drage og @friluftsliv.")
    assert res.text == (
        "ein [?knaskelure] og [?svibrent], [?pådrage] og [?friluftsliv]."
    )
    assert res.marks == [
        ("knaskelure", "unknown"),
        ("svibrent", "generation"),
        ("pådrage", "generation"),
        ("friluftsliv", "untranslated"),
    ]
    assert [res.text[o:o + 2] for o in res.offsets] == ["[?"] * 4


def test_parse_marks_leaves_no_raw_marks():
    res = core.parse_marks("ord# og *#to og @ alene *")
    for char in "*#@":
        assert char not in res.text


def test_hash_is_not_removed_by_u_in_fake(fake_apertium, tmp_path):
    # the fake behaves like the real program: -u strips only *
    import subprocess
    out = subprocess.run(
        [str(fake_apertium.bin_dir / "apertium"), "-u", "-f", "txt", "nob-nno"],
        input="svibrent xqztvik pådrage friluftsliv".encode(),
        capture_output=True, check=False,
    ).stdout.decode()
    assert out == "#svibrent xqztvik på#drage @friluftsliv"


def test_hash_marks_in_txt_are_marked_and_listed(fake_apertium, tmp_path, capsys):
    src = tmp_path / "hash.txt"
    src.write_text("Det er svibrent. Vi kan pådrage oss xqztvik.\n", encoding="utf-8")
    code, report, err = run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    assert code == 0, err
    text = (tmp_path / "ut" / "hash-nn.txt").read_text("utf-8")
    assert text == "Det er [?svibrent]. Vi kan [?pådrage] oss [?xqztvik].\n"
    found = {(u["word"], u["type"]) for u in report["outputs"][0]["unknown_words"]}
    assert found == {
        ("svibrent", "generation"),
        ("pådrage", "generation"),
        ("xqztvik", "unknown"),
    }


def test_authors_own_marks_survive_and_are_not_listed(fake_apertium, tmp_path, capsys):
    src = tmp_path / "eget.md"
    src.write_text(
        "# Tittel\n\n**held ope** @redaksjonen #sak *xqztvik* a#b\n", encoding="utf-8"
    )
    code, report, err = run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    assert code == 0, err
    text = (tmp_path / "ut" / "eget-nn.md").read_text("utf-8")
    assert text == "# Tittel\n\n**held ope** @redaksjonen #sak *[?xqztvik]* a#b\n"
    words = [u["word"] for u in report["outputs"][0]["unknown_words"]]
    assert words == ["xqztvik"]


def test_symbols_moved_by_apertium_cannot_damage_text(fake_apertium, tmp_path, capsys, monkeypatch):
    # the fake duplicates every symbol it receives, like the real transfer rules
    monkeypatch.setenv("FAKE_APERTIUM_MOVE_SYMBOLS", "1")
    src = tmp_path / "symboler.md"
    original = "# Tittel\n\nHan sa **17-åringen** var #2024 og @navn *ung* xqztvik.\n"
    src.write_text(original, encoding="utf-8")
    code, report, err = run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    assert code == 0, err
    text = (tmp_path / "ut" / "symboler-nn.md").read_text("utf-8")
    assert text == (
        "# Tittel\n\nHan sa **17-åringen** var #2024 og @navn *ung* [?xqztvik].\n"
    )
    assert [c for c in text if c in "*#@"] == [c for c in original if c in "*#@"]


def test_symbols_in_docx_survive_moving_fake(fake_apertium, tmp_path, capsys, monkeypatch):
    monkeypatch.setenv("FAKE_APERTIUM_MOVE_SYMBOLS", "1")
    src = tmp_path / "sym.docx"
    document = docx.Document()
    document.add_paragraph("Han sa **fra** skole")
    document.add_paragraph("#ærlig @navn")
    document.save(str(src))
    code, _, err = run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    assert code == 0, err
    paras = docx.Document(str(tmp_path / "ut" / "sym-nn.docx")).paragraphs
    assert [p.text for p in paras] == ["Han sa **frå** skule", "#ærlig @navn"]


def test_lost_line_structure_is_retried_line_by_line(fake_apertium, tmp_path, capsys, monkeypatch):
    monkeypatch.setenv("FAKE_APERTIUM_JOIN_LINES", "1")
    src = tmp_path / "linjer.txt"
    src.write_text("Fra skole\nfra måned\n\nfra flere\n", encoding="utf-8")
    code, report, err = run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    assert code == 0, err
    text = (tmp_path / "ut" / "linjer-nn.txt").read_text("utf-8")
    assert text == "Frå skule\nfrå månad\n\nfrå fleire\n"
    assert not any("ikke oversatt" in w for w in report["warnings"])


def test_line_that_cannot_be_translated_safely_is_kept_and_warned(
    fake_apertium, tmp_path, capsys, monkeypatch
):
    monkeypatch.setenv("FAKE_APERTIUM_SPLIT_WORD", "fleire")
    src = tmp_path / "linjer.txt"
    src.write_text("Fra skole\nfra flere\n\nfra måned\n", encoding="utf-8")
    code, report, err = run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    assert code == 0, err
    text = (tmp_path / "ut" / "linjer-nn.txt").read_text("utf-8")
    assert text == "Frå skule\nfra flere\n\nfrå månad\n"
    assert any("1 linjer ble ikke oversatt" in w for w in report["warnings"])


def test_docx_many_short_paragraphs_in_batches(fake_apertium, tmp_path, capsys):
    src = tmp_path / "mange.docx"
    document = docx.Document()
    for i in range(600):
        document.add_paragraph(f"Skole nummer {i}" + (" xqztvik" if i % 100 == 7 else ""))
    document.save(str(src))
    code, report, err = run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    assert code == 0, err
    paras = docx.Document(str(tmp_path / "ut" / "mange-nn.docx")).paragraphs
    assert [p.text.split(" ")[2] for p in paras] == [str(i) for i in range(600)]
    translate_calls = [c for c in fake_apertium.calls() if c != ["-l"]]
    assert len(translate_calls) > 1  # split up, not one call with 600 lines
    assert len(report["outputs"][0]["unknown_words"]) == 1
    assert report["outputs"][0]["unknown_words"][0]["count"] == 6


# --------------------------------------------------------------------------
# Text and Markdown
# --------------------------------------------------------------------------


def test_txt_to_nynorsk(fake_apertium, testdata, tmp_path, capsys):
    src = testdata / "avsnitt_nb.txt"
    before = sha(src)
    code, report, err = run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    assert code == 0, err
    assert sha(src) == before

    out = tmp_path / "ut" / "avsnitt_nb-nn.txt"
    assert report["outputs"][0]["path"] == str(out)
    text = out.read_text("utf-8")
    assert "frå neste månad" in text
    assert "[?fjordglimring]" in text
    assert "[?knaskelure]" in text
    assert "*" not in text
    # same paragraph structure and trailing newline as the original
    original = src.read_text("utf-8")
    assert text.count("\n\n") == original.count("\n\n") == 2
    assert text.endswith("\n") == original.endswith("\n")

    entry = report["outputs"][0]
    assert entry["draft"] is True
    assert entry["direction"] == "nob-nno"
    assert entry["paragraphs"] == 3
    words = {u["word"]: u for u in entry["unknown_words"]}
    assert words["fjordglimring"]["type"] == "unknown"
    assert words["knaskelure"]["type"] == "unknown"
    assert words["lånarane"]["type"] == "generation"
    assert words["friluftsliv"]["type"] == "untranslated"

    assert report["tool"] == "nynorsk"
    assert report["inputs"] == [{"path": str(src), "sha256": before}]
    assert report["summary"].startswith("Utkast:")
    assert any(w.startswith("UTKAST") for w in report["warnings"])
    assert any("fjordglimring" in w for w in report["warnings"])
    assert progress_lines(err)[-1] == 1.0


def test_apertium_called_with_argument_lists(fake_apertium, testdata, tmp_path, capsys):
    code, _, err = run(
        capsys, testdata / "avsnitt_nb.txt", "--to", "nn", "--out", tmp_path
    )
    assert code == 0, err
    calls = fake_apertium.calls()
    assert ["-l"] in calls
    assert ["-f", "txt", "nob-nno"] in calls
    assert all("-u" not in c for c in calls)


def test_direction_nb_uses_nno_nob(fake_apertium, testdata, tmp_path, capsys):
    code, report, err = run(
        capsys, testdata / "avsnitt_nn.txt", "--to", "nb", "--out", tmp_path
    )
    assert code == 0, err
    out = tmp_path / "avsnitt_nn-nb.txt"
    text = out.read_text("utf-8")
    assert "fra neste måned" in text
    assert report["outputs"][0]["direction"] == "nno-nob"
    assert "bokmål" in report["summary"]
    assert ["-f", "txt", "nno-nob"] in fake_apertium.calls()


def test_markdown_asterisks_are_not_unknown_words(fake_apertium, tmp_path, capsys):
    src = tmp_path / "notat.md"
    src.write_text(
        "# Biblioteket\n\nBiblioteket **holder åpent** i *Fiktivhavn*.\n\n- en knaskelure\n",
        encoding="utf-8",
    )
    code, report, err = run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    assert code == 0, err
    text = (tmp_path / "ut" / "notat-nn.md").read_text("utf-8")
    assert "**held ope**" in text
    assert "*[?Fiktivhavn]*" in text
    assert "- ein [?knaskelure]" in text
    words = [u["word"] for u in report["outputs"][0]["unknown_words"]]
    assert words == ["Fiktivhavn", "knaskelure"]


def test_odd_filename_and_multiple_inputs(fake_apertium, testdata, tmp_path, capsys):
    a = tmp_path / "inn" / "Søknad; ærlig åpen.txt"
    a.parent.mkdir()
    shutil.copy(testdata / "avsnitt_nb.txt", a)
    b = tmp_path / "inn" / "øya med mellomrom.docx"
    shutil.copy(testdata / "avsnitt_nb.docx", b)
    hashes = {a: sha(a), b: sha(b)}
    code, report, err = run(capsys, a, b, "--to", "nn", "--out", tmp_path / "ut ø")
    assert code == 0, err
    assert (tmp_path / "ut ø" / "Søknad; ærlig åpen-nn.txt").exists()
    assert (tmp_path / "ut ø" / "øya med mellomrom-nn.docx").exists()
    assert {p: sha(p) for p in hashes} == hashes
    assert len(report["inputs"]) == 2
    assert "2 filer" in report["summary"]


def test_never_overwrites(fake_apertium, testdata, tmp_path, capsys):
    src = testdata / "avsnitt_nb.txt"
    for _ in range(2):
        code, _, err = run(capsys, src, "--to", "nn", "--out", tmp_path)
        assert code == 0, err
    assert (tmp_path / "avsnitt_nb-nn.txt").exists()
    assert (tmp_path / "avsnitt_nb-nn-2.txt").exists()


def test_refuses_to_write_over_input(fake_apertium, tmp_path, capsys):
    src = tmp_path / "tekst-nn.txt"
    src.write_text("holder åpent\n", encoding="utf-8")
    # stem "tekst-nn" + "-nn.txt" never equals the input, but an input named
    # like the output in the same folder must stay untouched
    other = tmp_path / "tekst.txt"
    other.write_text("fra skole\n", encoding="utf-8")
    before = sha(src)
    code, report, err = run(capsys, other, "--to", "nn", "--out", tmp_path)
    assert code == 0, err
    assert sha(src) == before
    assert report["outputs"][0]["path"] == str(tmp_path / "tekst-nn-2.txt")


# --------------------------------------------------------------------------
# DOCX
# --------------------------------------------------------------------------


def _paragraphs(path):
    return [(p.style.name, p.text) for p in docx.Document(str(path)).paragraphs]


def test_docx_keeps_order_and_paragraph_styles(fake_apertium, testdata, tmp_path, capsys):
    src = testdata / "avsnitt_nb.docx"
    before = sha(src)
    code, report, err = run(capsys, src, "--to", "nn", "--out", tmp_path)
    assert code == 0, err
    assert sha(src) == before
    out = tmp_path / "avsnitt_nb-nn.docx"

    original = _paragraphs(src)
    result = _paragraphs(out)
    assert [s for s, _ in result] == [s for s, _ in original]
    assert result[0] == ("Heading 1", "Lengre opningstider på biblioteket")
    bullets = [t for s, t in result if s == "List Bullet"]
    assert len(bullets) == 3
    assert "[?fjordglimring]" in bullets[1]
    assert "[?knaskelure]" in bullets[2]
    # every paragraph with text was translated, empty ones stay empty
    for (_, old), (_, new) in zip(original, result):
        assert bool(old.strip()) == bool(new.strip())
    assert any("held ope lenger om kvelden" in t for _, t in result)

    entry = report["outputs"][0]
    assert entry["draft"] is True
    assert entry["paragraphs"] == sum(1 for _, t in original if t.strip())
    words = [u["word"] for u in entry["unknown_words"]]
    assert "fjordglimring" in words and "knaskelure" in words
    # the bold phrase is run-level formatting, which is reported as lost
    assert any("blandet tegnformatering" in w for w in report["warnings"])


def _make_docx(path):
    document = docx.Document()
    p = document.add_paragraph("Biblioteket holder åpent", style="Heading 2")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = document.add_paragraph("Linje en fra skole")
    p.add_run().add_break()
    p.add_run("linje to med en knaskelure")
    document.add_paragraph("")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "flere barn"
    table.cell(0, 1).text = "Fiktivhavn"
    document.add_paragraph("Slutt", style="Quote")
    document.save(str(path))


@pytest.mark.parametrize("join_lines", [False, True])
def test_docx_line_breaks_tables_and_fallback(
    fake_apertium, tmp_path, capsys, monkeypatch, join_lines
):
    if join_lines:
        # Apertium loses the newlines: the tool must fall back to one call
        # per paragraph and still get the paragraphs right
        monkeypatch.setenv("FAKE_APERTIUM_JOIN_LINES", "1")
    src = tmp_path / "eget dokument.docx"
    _make_docx(src)
    code, report, err = run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    assert code == 0, err
    document = docx.Document(str(tmp_path / "ut" / "eget dokument-nn.docx"))
    paras = document.paragraphs
    assert [p.style.name for p in paras] == ["Heading 2", "Normal", "Normal", "Quote"]
    assert paras[0].text == "Biblioteket held ope"
    assert paras[0].alignment == WD_ALIGN_PARAGRAPH.CENTER
    if not join_lines:
        assert paras[1].text == "Linje ein frå skule\nlinje to med ein [?knaskelure]"
    else:
        assert "[?knaskelure]" in paras[1].text
    assert paras[2].text == ""
    cells = [c.text for c in document.tables[0].rows[0].cells]
    assert cells == ["fleire born", "[?Fiktivhavn]"]
    words = [u["word"] for u in report["outputs"][0]["unknown_words"]]
    assert words == ["knaskelure", "Fiktivhavn"]
    calls = fake_apertium.calls()
    translate_calls = [c for c in calls if c != ["-l"]]
    assert len(translate_calls) == (1 if not join_lines else 8)


def test_docx_keeps_page_breaks(fake_apertium, tmp_path, capsys):
    from docx.enum.text import WD_BREAK

    src = tmp_path / "bilde.docx"
    document = docx.Document()
    p = document.add_paragraph("holder åpent")
    p.add_run().add_break(WD_BREAK.PAGE)
    p.add_run(" fra skole")
    document.save(str(src))
    code, _, err = run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    assert code == 0, err
    out = docx.Document(str(tmp_path / "ut" / "bilde-nn.docx"))
    xml = out.paragraphs[0]._p.xml
    assert 'w:type="page"' in xml
    assert out.paragraphs[0].text.replace("\n", "").strip().startswith("held ope")


# --------------------------------------------------------------------------
# Errors: Norwegian message, code 1, no stack trace
# --------------------------------------------------------------------------


def _assert_user_error(code, err, *needles):
    assert code == 1
    assert err.startswith("Feil:") or "\nFeil:" in err
    assert "Traceback" not in err
    for needle in needles:
        assert needle in err


def test_missing_apertium(no_apertium, testdata, tmp_path, capsys):
    code, _, err = run(capsys, testdata / "avsnitt_nb.txt", "--to", "nn", "--out", tmp_path)
    _assert_user_error(code, err, "Apertium er ikke installert", "OPPSETT.md", "apertium-nno-nob")
    assert not list(tmp_path.glob("*-nn.*"))


def test_missing_language_pair_from_list(fake_apertium, testdata, tmp_path, capsys, monkeypatch):
    monkeypatch.setenv("FAKE_APERTIUM_MODES", "nob-nno")
    code, _, err = run(capsys, testdata / "avsnitt_nn.txt", "--to", "nb", "--out", tmp_path)
    _assert_user_error(code, err, "«nno-nob» finnes ikke", "OPPSETT.md")


def test_missing_language_pair_without_list(fake_apertium, testdata, tmp_path, capsys, monkeypatch):
    monkeypatch.setenv("FAKE_APERTIUM_NO_LIST", "1")
    monkeypatch.setenv("FAKE_APERTIUM_MODES", "eng-spa")
    code, _, err = run(capsys, testdata / "avsnitt_nb.txt", "--to", "nn", "--out", tmp_path)
    _assert_user_error(code, err, "«nob-nno» finnes ikke")


def test_apertium_crash(fake_apertium, testdata, tmp_path, capsys, monkeypatch):
    monkeypatch.setenv("FAKE_APERTIUM_FAIL", "1")
    code, _, err = run(capsys, testdata / "avsnitt_nb.txt", "--to", "nn", "--out", tmp_path)
    _assert_user_error(code, err, "Apertium feilet med kode 3", "Segmentation fault")


def test_file_not_found(fake_apertium, tmp_path, capsys):
    code, _, err = run(capsys, tmp_path / "finnes ikke.txt", "--to", "nn", "--out", tmp_path)
    _assert_user_error(code, err, "finnes ikke")


def test_empty_file(fake_apertium, tmp_path, capsys):
    src = tmp_path / "tom.txt"
    src.write_text("", encoding="utf-8")
    code, _, err = run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    _assert_user_error(code, err, "er tom")


def test_whitespace_only_file(fake_apertium, tmp_path, capsys):
    src = tmp_path / "blank.md"
    src.write_text("\n  \n", encoding="utf-8")
    code, _, err = run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    _assert_user_error(code, err, "ingen tekst")


def test_wrong_format(fake_apertium, tmp_path, capsys):
    src = tmp_path / "dokument.pdf"
    src.write_bytes(b"%PDF-1.4\n")
    code, _, err = run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    _assert_user_error(code, err, "ikke støttes", ".docx")


def test_fake_docx(fake_apertium, tmp_path, capsys):
    src = tmp_path / "falsk.docx"
    src.write_text("dette er ikke et Word-dokument", encoding="utf-8")
    code, _, err = run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    _assert_user_error(code, err, "ikke en gyldig DOCX-fil")
    assert not (tmp_path / "ut").exists() or not any((tmp_path / "ut").iterdir())


def test_fake_docx_reported_before_apertium_is_checked(no_apertium, tmp_path, capsys):
    src = tmp_path / "falsk.docx"
    src.write_text("dette er ikke et Word-dokument", encoding="utf-8")
    code, _, err = run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    _assert_user_error(code, err, "ikke en gyldig DOCX-fil")
    assert "Apertium er ikke installert" not in err


def test_docx_without_text(fake_apertium, tmp_path, capsys):
    src = tmp_path / "tomt.docx"
    docx.Document().save(str(src))
    code, _, err = run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    _assert_user_error(code, err, "ingen tekst")


def test_not_utf8(fake_apertium, tmp_path, capsys):
    src = tmp_path / "latin1.txt"
    src.write_bytes("Søknad om støtte".encode("latin-1"))
    code, _, err = run(capsys, src, "--to", "nn", "--out", tmp_path / "ut")
    _assert_user_error(code, err, "UTF-8")


def test_missing_direction(fake_apertium, testdata, tmp_path, capsys):
    code, _, err = run(capsys, testdata / "avsnitt_nb.txt", "--out", tmp_path)
    _assert_user_error(code, err, "Ugyldige argumenter")


def test_bad_input_checked_before_apertium(fake_apertium, testdata, tmp_path, capsys):
    code, _, err = run(
        capsys, testdata / "avsnitt_nb.txt", tmp_path / "mangler.md",
        "--to", "nn", "--out", tmp_path / "ut",
    )
    _assert_user_error(code, err, "mangler.md")
    assert fake_apertium.calls() == []


# --------------------------------------------------------------------------
# Registration
# --------------------------------------------------------------------------


def test_registered_with_jv(fake_apertium, testdata, tmp_path, capsys):
    assert "nynorsk" in jv_cli.discover_tools()
    code = jv_cli.main(
        ["nynorsk", str(testdata / "avsnitt_nb.txt"), "--to", "nn", "--out", str(tmp_path)]
    )
    assert code == 0
    report = json.loads(capsys.readouterr().out)
    assert report["tool"] == "nynorsk"


def test_help_is_norwegian(capsys):
    with pytest.raises(SystemExit):
        cli._build_parser().parse_args(["--help"])
    out = capsys.readouterr().out
    assert "målform" in out and "Originalfilen endres aldri" in out

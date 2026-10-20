"""Core logic for compare: read two documents, diff paragraphs, then words.

No printing and no sys.exit here. Errors the user can fix raise UserError.
"""

from __future__ import annotations

import difflib
import html
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from jv_common import UserError

SUPPORTED_SUFFIXES = (".txt", ".md", ".docx", ".pdf")
SCANNED_WARNING = "PDF-en ser ut til å være skannet. Kjør `ocr` først."

# Minimum number of words for a sentence to be detected as moved.
MIN_MOVED_WORDS = 4
# Minimum similarity for two paragraphs in a replaced block to be paired.
PAIR_THRESHOLD = 0.4

# Numbers with Norwegian thousands separators ("1 200 000", "1 200,50") and
# decimals ("3,5") are one token, so a changed amount reads as one change.
_TOKEN_RE = re.compile(
    r"\d{1,3}(?:[   ]\d{3})+(?:,\d+)?"
    r"|\d+(?:[.,]\d+)+"
    r"|\w+(?:[-'’]\w+)*"
    r"|[^\w\s]"
)
_WS_RE = re.compile(r"\s+")
_PARA_SPLIT_RE = re.compile(r"\n[ \t\r\f\v]*\n\s*")
# A sentence ends with . ! ? or … (optionally followed by closing quotes or
# brackets), then whitespace, then something that looks like a sentence start.
_SENTENCE_END_RE = re.compile(
    r"[.!?…]+[\"»”’)\]]*(?=\s+[\"«“(\[]?[A-ZÆØÅÄÖÜÉ])"
)
_END_PUNCT = tuple(".!?:;…»\"”)")


# --------------------------------------------------------------------------
# Reading documents
# --------------------------------------------------------------------------


@dataclass
class Document:
    path: Path
    paragraphs: list[str]
    warnings: list[str] = field(default_factory=list)


def read_document(path: str | Path) -> Document:
    """Read a supported file into a list of paragraphs (raw text)."""
    p = Path(path)
    if not p.exists():
        raise UserError(f"Finner ikke filen «{p}».")
    if not p.is_file():
        raise UserError(f"«{p}» er ikke en fil.")
    suffix = p.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UserError(
            f"Filtypen «{p.suffix or '(ingen)'}» støttes ikke for «{p.name}». "
            "Bruk .txt, .md, .docx eller .pdf."
        )
    if p.stat().st_size == 0:
        raise UserError(f"Filen «{p.name}» er tom.")
    if suffix in (".txt", ".md"):
        doc = _read_text(p)
    elif suffix == ".docx":
        doc = _read_docx(p)
    else:
        doc = _read_pdf(p)
    if not doc.paragraphs and not doc.warnings:
        raise UserError(f"Fant ingen tekst i «{p.name}».")
    return doc


def _read_text(p: Path) -> Document:
    data = p.read_bytes()
    if b"\x00" in data:
        raise UserError(f"«{p.name}» ser ikke ut til å være en tekstfil.")
    warnings: list[str] = []
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
        warnings.append(
            f"{p.name}: filen er ikke UTF-8 og er lest som Windows-1252. "
            "Sjekk at æ, ø og å ser riktige ut."
        )
    return Document(p, split_paragraphs(text), warnings)


def split_paragraphs(text: str) -> list[str]:
    """Split plain text into paragraphs at blank lines."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    parts = _PARA_SPLIT_RE.split(text)
    return [part.strip("\n") for part in parts if part.strip()]


def _read_docx(p: Path) -> Document:
    try:
        import docx  # python-docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        document = docx.Document(str(p))
    except (zipfile.BadZipFile, KeyError, ValueError) as exc:
        raise UserError(f"«{p.name}» kunne ikke leses som DOCX-fil.") from exc
    except Exception as exc:  # python-docx raises several types on bad files
        if type(exc).__name__ in ("PackageNotFoundError", "XMLSyntaxError"):
            raise UserError(f"«{p.name}» kunne ikke leses som DOCX-fil.") from exc
        raise

    paragraphs: list[str] = []
    body = document.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            text = Paragraph(child, document).text
            if text.strip():
                paragraphs.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            for row in table.rows:
                cells: list[str] = []
                seen: set[int] = set()
                for cell in row.cells:
                    # merged cells appear several times in python-docx
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    cells.append(" ".join(cell.text.split()))
                line = " | ".join(cells)
                if line.strip(" |"):
                    paragraphs.append(line)
    return Document(p, paragraphs)


def _join_lines(lines: list[str]) -> str:
    out = ""
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if not out:
            out = line
        elif out.endswith("­"):
            out = out[:-1] + line
        elif out.endswith("-") and line[:1].islower() and out[-2:-1].isalpha():
            # hyphenated word at line break: "kultur-" + "huset"
            out = out[:-1] + line
        else:
            out = out + " " + line
    return out


def _read_pdf(p: Path) -> Document:
    import pymupdf

    try:
        pdf = pymupdf.open(str(p))
    except Exception as exc:  # pymupdf raises its own FileDataError etc.
        raise UserError(f"«{p.name}» kunne ikke leses som PDF-fil.") from exc
    with pdf:
        if pdf.needs_pass:
            raise UserError(f"«{p.name}» er passordbeskyttet.")
        if not pdf.is_pdf:
            raise UserError(f"«{p.name}» er ikke en PDF-fil.")
        paragraphs: list[str] = []
        pages_without_text: list[int] = []
        for index, page in enumerate(pdf, start=1):
            blocks = [
                b for b in page.get_text("blocks", sort=True) if b[6] == 0
            ]
            texts = [_join_lines(b[4].split("\n")) for b in blocks]
            texts = [t for t in texts if t]
            if not texts:
                pages_without_text.append(index)
                continue
            # paragraph continued over a page break
            if (
                paragraphs
                and index > 1
                and not paragraphs[-1].endswith(_END_PUNCT)
                and texts[0][:1].islower()
            ):
                paragraphs[-1] = paragraphs[-1] + " " + texts.pop(0)
            paragraphs.extend(texts)
        page_count = pdf.page_count

    warnings: list[str] = []
    if pages_without_text:
        if len(pages_without_text) == page_count:
            warnings.append(f"{p.name}: {SCANNED_WARNING}")
        else:
            pages = ", ".join(str(n) for n in pages_without_text)
            warnings.append(
                f"{p.name}: side {pages} mangler tekstlag. {SCANNED_WARNING}"
            )
    return Document(p, paragraphs, warnings)


# --------------------------------------------------------------------------
# Tokens and sentences
# --------------------------------------------------------------------------


@dataclass
class Token:
    text: str  # the word or punctuation; empty for a leading-space token
    space: str  # whitespace that follows the token
    key: str  # what is compared
    moved: int | None = None  # index into Comparison.moved for a moved sentence

    @property
    def is_word(self) -> bool:
        return self.moved is None and any(ch.isalnum() for ch in self.text)


@dataclass
class Options:
    ignore_whitespace: bool = True
    ignore_case: bool = False


def _fold(text: str, opts: Options) -> str:
    return text.casefold() if opts.ignore_case else text


def tokenize(text: str, opts: Options) -> list[Token]:
    tokens: list[Token] = []
    pos = 0
    for match in _TOKEN_RE.finditer(text):
        gap = text[pos : match.start()]
        if gap:
            if tokens:
                tokens[-1].space += gap
            elif not opts.ignore_whitespace:
                tokens.append(Token("", gap, gap))
        word = match.group()
        if opts.ignore_whitespace:
            word = _WS_RE.sub(" ", word)
        tokens.append(Token(word, "", _fold(word, opts)))
        pos = match.end()
    tail = text[pos:]
    if tokens and tail:
        tokens[-1].space += tail
    for tok in tokens:
        if opts.ignore_whitespace:
            tok.space = " " if tok.space else ""
        else:
            tok.key = tok.key + "\x00" + tok.space
    if opts.ignore_whitespace and tokens:
        tokens[-1].space = ""
    return tokens


def split_sentences(text: str) -> list[tuple[int, int]]:
    """Return (start, end) spans of the sentences in a paragraph."""
    spans: list[tuple[int, int]] = []
    start = 0
    for match in _SENTENCE_END_RE.finditer(text):
        spans.append((start, match.end()))
        start = match.end()
    spans.append((start, len(text)))
    result = []
    for s, e in spans:
        while s < e and text[s].isspace():
            s += 1
        if s < e:
            result.append((s, e))
    return result


def _sentence_key(text: str, opts: Options) -> str:
    # moves are detected regardless of whitespace, so reflowed text still matches
    return _fold(_WS_RE.sub(" ", text).strip(), opts)


def _word_count(text: str) -> int:
    return sum(1 for t in _TOKEN_RE.findall(text) if any(c.isalnum() for c in t))


# --------------------------------------------------------------------------
# Comparison
# --------------------------------------------------------------------------


@dataclass
class Moved:
    text: str
    old_paragraph: int  # 1-based paragraph number in the old version
    new_paragraph: int  # 1-based paragraph number in the new version


@dataclass
class Row:
    """One line in the two-column report."""

    kind: str  # "equal", "changed", "deleted", "inserted"
    old_no: int | None
    new_no: int | None
    old_tokens: list[tuple[str, Token]]  # (op, token), op in equal/delete
    new_tokens: list[tuple[str, Token]]  # (op, token), op in equal/insert


@dataclass
class Comparison:
    old: Document
    new: Document
    options: Options
    rows: list[Row]
    moved: list[Moved]
    changed_paragraphs: int
    deleted_paragraphs: int
    inserted_paragraphs: int
    words_added: int
    words_removed: int

    @property
    def total_changed(self) -> int:
        return (
            self.changed_paragraphs
            + self.deleted_paragraphs
            + self.inserted_paragraphs
        )

    @property
    def warnings(self) -> list[str]:
        return self.old.warnings + self.new.warnings


def _para_key(text: str, opts: Options) -> str:
    if opts.ignore_whitespace:
        text = _WS_RE.sub(" ", text).strip()
    return _fold(text, opts)


def _similarity(a: list[str], b: list[str]) -> float:
    matcher = difflib.SequenceMatcher(None, a, b, autojunk=False)
    if matcher.real_quick_ratio() < PAIR_THRESHOLD:
        return 0.0
    if matcher.quick_ratio() < PAIR_THRESHOLD:
        return 0.0
    return matcher.ratio()


def _pair_block(
    old_idx: list[int],
    new_idx: list[int],
    old_words: list[list[str]],
    new_words: list[list[str]],
) -> list[tuple[int | None, int | None]]:
    """Pair paragraphs inside a replaced block, keeping order.

    Dynamic programming that maximises the summed similarity of pairs above
    PAIR_THRESHOLD. Unpaired paragraphs become deletions or insertions.
    """
    n, m = len(old_idx), len(new_idx)
    sim = [
        [_similarity(old_words[i], new_words[j]) for j in new_idx] for i in old_idx
    ]
    best = [[0.0] * (m + 1) for _ in range(n + 1)]
    for i in range(n - 1, -1, -1):
        for j in range(m - 1, -1, -1):
            score = best[i + 1][j]
            score = max(score, best[i][j + 1])
            if sim[i][j] >= PAIR_THRESHOLD:
                score = max(score, sim[i][j] + best[i + 1][j + 1])
            best[i][j] = score
    pairs: list[tuple[int | None, int | None]] = []
    i = j = 0
    while i < n and j < m:
        if sim[i][j] >= PAIR_THRESHOLD and best[i][j] == sim[i][j] + best[i + 1][j + 1]:
            pairs.append((old_idx[i], new_idx[j]))
            i += 1
            j += 1
        elif best[i][j] == best[i + 1][j]:
            pairs.append((old_idx[i], None))
            i += 1
        else:
            pairs.append((None, new_idx[j]))
            j += 1
    pairs.extend((old_idx[k], None) for k in range(i, n))
    pairs.extend((None, new_idx[k]) for k in range(j, m))
    return pairs


def align_paragraphs(
    old: list[str], new: list[str], opts: Options
) -> list[tuple[int | None, int | None]]:
    """Align paragraphs: difflib first, then pair similar ones in replaced blocks."""
    old_keys = [_para_key(t, opts) for t in old]
    new_keys = [_para_key(t, opts) for t in new]
    old_words = [[t.key for t in tokenize(x, opts)] for x in old]
    new_words = [[t.key for t in tokenize(x, opts)] for x in new]
    matcher = difflib.SequenceMatcher(None, old_keys, new_keys, autojunk=False)
    pairs: list[tuple[int | None, int | None]] = []
    for op, i1, i2, j1, j2 in matcher.get_opcodes():
        if op == "equal":
            pairs.extend(zip(range(i1, i2), range(j1, j2)))
        elif op == "delete":
            pairs.extend((i, None) for i in range(i1, i2))
        elif op == "insert":
            pairs.extend((None, j) for j in range(j1, j2))
        else:
            pairs.extend(
                _pair_block(
                    list(range(i1, i2)), list(range(j1, j2)), old_words, new_words
                )
            )
    return pairs


def _find_moved(
    old: list[str],
    new: list[str],
    pairs: list[tuple[int | None, int | None]],
    opts: Options,
) -> tuple[list[Moved], dict[tuple[int, int], int], dict[tuple[int, int], int]]:
    """Find sentences that left one paragraph and turned up in another.

    Returns the moves and maps (paragraph, sentence) -> move index for the
    old and the new side.
    """
    old_partner = {i: j for i, j in pairs if i is not None}
    new_partner = {j: i for i, j in pairs if j is not None}

    def sentences(texts: list[str]) -> list[list[tuple[str, str]]]:
        out = []
        for text in texts:
            out.append(
                [
                    (_sentence_key(text[s:e], opts), text[s:e])
                    for s, e in split_sentences(text)
                ]
            )
        return out

    old_sent = sentences(old)
    new_sent = sentences(new)
    old_keys = [{k for k, _ in para} for para in old_sent]
    new_keys = [{k for k, _ in para} for para in new_sent]

    new_locations: dict[str, list[tuple[int, int]]] = {}
    for j, para in enumerate(new_sent):
        for s, (key, _) in enumerate(para):
            new_locations.setdefault(key, []).append((j, s))

    moved: list[Moved] = []
    old_map: dict[tuple[int, int], int] = {}
    new_map: dict[tuple[int, int], int] = {}
    used_new: set[tuple[int, int]] = set()
    for i, para in enumerate(old_sent):
        partner = old_partner.get(i)
        for s, (key, text) in enumerate(para):
            if _word_count(text) < MIN_MOVED_WORDS:
                continue
            if partner is not None and key in new_keys[partner]:
                continue  # still in the same paragraph
            for j, t in new_locations.get(key, []):
                if (j, t) in used_new or j == partner:
                    continue
                back = new_partner.get(j)
                if back is not None and key in old_keys[back]:
                    continue  # the new paragraph already had it
                used_new.add((j, t))
                old_map[(i, s)] = len(moved)
                new_map[(j, t)] = len(moved)
                moved.append(Moved(_WS_RE.sub(" ", text).strip(), i + 1, j + 1))
                break
    return moved, old_map, new_map


def _tokens_with_moves(
    text: str, para_no: int, move_map: dict[tuple[int, int], int], opts: Options
) -> list[Token]:
    """Tokenize a paragraph; each moved sentence becomes one marker token."""
    spans = split_sentences(text)
    if not any((para_no, s) in move_map for s in range(len(spans))):
        return tokenize(text, opts)
    tokens: list[Token] = []
    pos = 0
    for s, (start, end) in enumerate(spans):
        if (para_no, s) not in move_map:
            continue
        before = tokenize(text[pos:start], opts)
        tokens.extend(before)
        idx = move_map[(para_no, s)]
        sentence = text[start:end]
        rest = text[end:]
        space = rest[: len(rest) - len(rest.lstrip())]
        if opts.ignore_whitespace:
            space = " " if space else ""
            sentence = _WS_RE.sub(" ", sentence).strip()
        if tokens and opts.ignore_whitespace and not tokens[-1].space:
            tokens[-1].space = " "
        tokens.append(Token(sentence, space, f"\x01moved{idx}", moved=idx))
        pos = end + len(text[end:]) - len(text[end:].lstrip())
    tokens.extend(tokenize(text[pos:], opts))
    if opts.ignore_whitespace and tokens:
        tokens[-1].space = ""
    return tokens


def compare_documents(old: Document, new: Document, opts: Options) -> Comparison:
    pairs = align_paragraphs(old.paragraphs, new.paragraphs, opts)
    moved, old_map, new_map = _find_moved(old.paragraphs, new.paragraphs, pairs, opts)

    rows: list[Row] = []
    changed = deleted = inserted = added = removed = 0
    for i, j in pairs:
        old_toks = (
            _tokens_with_moves(old.paragraphs[i], i, old_map, opts)
            if i is not None
            else []
        )
        new_toks = (
            _tokens_with_moves(new.paragraphs[j], j, new_map, opts)
            if j is not None
            else []
        )
        if i is not None and j is None:
            deleted += 1
            removed += sum(t.is_word for t in old_toks)
            rows.append(Row("deleted", i + 1, None, [("delete", t) for t in old_toks], []))
            continue
        if i is None and j is not None:
            inserted += 1
            added += sum(t.is_word for t in new_toks)
            rows.append(Row("inserted", None, j + 1, [], [("insert", t) for t in new_toks]))
            continue
        assert i is not None and j is not None
        matcher = difflib.SequenceMatcher(
            None, [t.key for t in old_toks], [t.key for t in new_toks], autojunk=False
        )
        left: list[tuple[str, Token]] = []
        right: list[tuple[str, Token]] = []
        is_changed = False
        for op, i1, i2, j1, j2 in matcher.get_opcodes():
            if op == "equal":
                left.extend(("equal", t) for t in old_toks[i1:i2])
                right.extend(("equal", t) for t in new_toks[j1:j2])
                continue
            is_changed = True
            left.extend(("delete", t) for t in old_toks[i1:i2])
            right.extend(("insert", t) for t in new_toks[j1:j2])
            removed += sum(t.is_word for t in old_toks[i1:i2])
            added += sum(t.is_word for t in new_toks[j1:j2])
        if is_changed:
            changed += 1
        rows.append(Row("changed" if is_changed else "equal", i + 1, j + 1, left, right))

    return Comparison(
        old=old,
        new=new,
        options=opts,
        rows=rows,
        moved=moved,
        changed_paragraphs=changed,
        deleted_paragraphs=deleted,
        inserted_paragraphs=inserted,
        words_added=added,
        words_removed=removed,
    )


def compare_files(
    old_path: str | Path, new_path: str | Path, opts: Options | None = None
) -> Comparison:
    opts = opts or Options()
    return compare_documents(read_document(old_path), read_document(new_path), opts)


# --------------------------------------------------------------------------
# Norwegian wording
# --------------------------------------------------------------------------


def _n(count: int, singular: str, plural: str) -> str:
    return f"{count} {singular if count == 1 else plural}"


def summary_text(c: Comparison) -> str:
    if c.total_changed == 0:
        return "Ingen endringer funnet mellom de to versjonene."
    parts = [
        f"{_n(c.total_changed, 'avsnitt', 'avsnitt')} er endret "
        f"({c.changed_paragraphs} endret, {c.deleted_paragraphs} slettet, "
        f"{c.inserted_paragraphs} nye)",
        f"{_n(c.words_added, 'ord', 'ord')} lagt til og "
        f"{_n(c.words_removed, 'ord', 'ord')} fjernet",
    ]
    text = ", ".join(parts)
    if c.moved:
        text += f", {_n(len(c.moved), 'setning', 'setninger')} flyttet"
    return text + "."


# --------------------------------------------------------------------------
# HTML report
# --------------------------------------------------------------------------

_CSS = """
:root { color-scheme: light dark; --bg:#ffffff; --fg:#1d1d1f; --muted:#6b6b70;
  --line:#d9d9de; --head:#f4f4f6; --del-bg:#fde2e1; --del-fg:#a1140c;
  --ins-bg:#dcf5e3; --ins-fg:#0b6b2c; --mov:#5b3fb5; --row-changed:#fffbea; }
@media (prefers-color-scheme: dark) {
  :root { --bg:#17171a; --fg:#ececf0; --muted:#a0a0a8; --line:#3a3a40;
    --head:#222227; --del-bg:#4a1d1c; --del-fg:#ffb3ad; --ins-bg:#16391f;
    --ins-fg:#9be5ad; --mov:#c6b5ff; --row-changed:#2a2618; }
}
* { box-sizing: border-box; }
body { margin: 0; padding: 24px 16px 48px; background: var(--bg); color: var(--fg);
  font: 16px/1.55 -apple-system, BlinkMacSystemFont, "Segoe UI", Helvetica, Arial, sans-serif; }
main { max-width: 1200px; margin: 0 auto; }
h1 { font-size: 1.5rem; margin: 0 0 4px; }
.files { color: var(--muted); margin: 0 0 20px; overflow-wrap: anywhere; }
.stats { display: flex; flex-wrap: wrap; gap: 12px; margin: 0 0 16px; padding: 0; list-style: none; }
.stats li { border: 1px solid var(--line); border-radius: 8px; padding: 10px 14px; min-width: 140px; }
.stats b { display: block; font-size: 1.6rem; line-height: 1.2; font-variant-numeric: tabular-nums; }
.stats span { color: var(--muted); font-size: .9rem; }
.legend, .options { color: var(--muted); font-size: .9rem; margin: 0 0 8px; }
.warnings { border: 1px solid var(--del-fg); border-radius: 8px; padding: 10px 14px; margin: 0 0 16px; }
.warnings p { margin: 4px 0; }
.moved-list { margin: 0 0 16px; }
.moved-list li { margin: 2px 0; }
table { width: 100%; border-collapse: collapse; table-layout: fixed; }
th, td { border: 1px solid var(--line); padding: 8px 10px; vertical-align: top; text-align: left; }
th { background: var(--head); }
td { white-space: pre-wrap; overflow-wrap: anywhere; }
th.no, td.no { width: 3.2em; color: var(--muted); font-size: .85rem; text-align: right;
  font-variant-numeric: tabular-nums; }
tr.changed td, tr.deleted td, tr.inserted td { background: var(--row-changed); }
tr.equal td.text { color: var(--muted); }
del { background: var(--del-bg); color: var(--del-fg); text-decoration: line-through; }
ins { background: var(--ins-bg); color: var(--ins-fg); text-decoration: none; }
del.moved, ins.moved { outline: 2px dashed var(--mov); outline-offset: 1px; }
.tag { font-size: .75rem; color: var(--mov); font-weight: 600; white-space: nowrap; }
.empty { color: var(--muted); font-style: italic; }
@media (max-width: 640px) {
  body { font-size: 15px; }
  th, td { padding: 6px; }
  th.no, td.no { width: 2.4em; }
}
@media print { tr { break-inside: avoid; } }
"""


def _esc(text: str) -> str:
    return html.escape(text, quote=True)


def _render_side(items: list[tuple[str, Token]], comp: Comparison, side: str) -> str:
    out: list[str] = []
    run_op: str | None = None
    run: list[str] = []
    run_space = ""

    def flush() -> None:
        nonlocal run_op, run, run_space
        if run_op is None:
            return
        body = "".join(run)
        if run_op == "equal":
            out.append(body)
        else:
            tag = "del" if run_op == "delete" else "ins"
            out.append(f"<{tag}>{body}</{tag}>")
        out.append(_esc(run_space))
        run_op, run, run_space = None, [], ""

    for op, tok in items:
        if tok.moved is not None:
            flush()
            mv = comp.moved[tok.moved]
            tag = "del" if side == "old" else "ins"
            label = (
                f"flyttet til avsnitt {mv.new_paragraph}"
                if side == "old"
                else f"flyttet fra avsnitt {mv.old_paragraph}"
            )
            out.append(
                f'<{tag} class="moved" title="{_esc(label)}">{_esc(tok.text)}</{tag}>'
                f' <span class="tag">[{_esc(label)}]</span>{_esc(tok.space)}'
            )
            continue
        if op != run_op:
            flush()
            run_op = op
        else:
            run.append(_esc(run_space))
        run.append(_esc(tok.text))
        run_space = tok.space
    flush()
    return "".join(out)


def render_html(c: Comparison) -> str:
    old_name = c.old.path.name
    new_name = c.new.path.name
    opts = []
    opts.append(
        "mellomrom ignoreres" if c.options.ignore_whitespace else "mellomrom telles med"
    )
    opts.append(
        "store og små bokstaver regnes som like"
        if c.options.ignore_case
        else "store og små bokstaver skilles"
    )

    parts: list[str] = [
        "<!doctype html>",
        '<html lang="no">',
        "<head>",
        '<meta charset="utf-8">',
        '<meta name="viewport" content="width=device-width, initial-scale=1">',
        f"<title>Sammenligning: {_esc(old_name)} mot {_esc(new_name)}</title>",
        f"<style>{_CSS}</style>",
        "</head>",
        "<body>",
        "<main>",
        "<h1>Sammenligning av to versjoner</h1>",
        f'<p class="files">Gammel: <strong>{_esc(old_name)}</strong> '
        f"({len(c.old.paragraphs)} avsnitt) &nbsp;→&nbsp; Ny: "
        f"<strong>{_esc(new_name)}</strong> ({len(c.new.paragraphs)} avsnitt)</p>",
        '<ul class="stats">',
        f"<li><b>{c.total_changed}</b><span>endrede avsnitt "
        f"({c.changed_paragraphs} endret, {c.deleted_paragraphs} slettet, "
        f"{c.inserted_paragraphs} nye)</span></li>",
        f"<li><b>{c.words_added}</b><span>ord lagt til</span></li>",
        f"<li><b>{c.words_removed}</b><span>ord fjernet</span></li>",
        f"<li><b>{len(c.moved)}</b><span>flyttede setninger</span></li>",
        "</ul>",
        '<p class="legend"><del>Slettet tekst</del> &nbsp; <ins>Ny tekst</ins> &nbsp; '
        '<del class="moved">Flyttet</del>/<ins class="moved">flyttet</ins> setning '
        "(telles ikke som lagt til eller fjernet)</p>",
        f'<p class="options">Valg: {_esc(", ".join(opts))}.</p>',
    ]
    if c.warnings:
        parts.append('<div class="warnings" role="alert">')
        parts.extend(f"<p>{_esc(w)}</p>" for w in c.warnings)
        parts.append("</div>")
    if c.moved:
        parts.append('<ol class="moved-list">')
        for mv in c.moved:
            parts.append(
                f"<li>«{_esc(mv.text)}» flyttet fra avsnitt {mv.old_paragraph} "
                f"til avsnitt {mv.new_paragraph}</li>"
            )
        parts.append("</ol>")
    parts.append("<table>")
    parts.append(
        '<thead><tr><th class="no">Nr.</th>'
        f"<th>Gammel versjon: {_esc(old_name)}</th>"
        '<th class="no">Nr.</th>'
        f"<th>Ny versjon: {_esc(new_name)}</th></tr></thead>"
    )
    parts.append("<tbody>")
    for row in c.rows:
        left = (
            _render_side(row.old_tokens, c, "old")
            if row.old_no is not None
            else '<span class="empty">(nytt avsnitt)</span>'
        )
        right = (
            _render_side(row.new_tokens, c, "new")
            if row.new_no is not None
            else '<span class="empty">(avsnittet er slettet)</span>'
        )
        old_no = "" if row.old_no is None else str(row.old_no)
        new_no = "" if row.new_no is None else str(row.new_no)
        parts.append(
            f'<tr class="{row.kind}"><td class="no">{old_no}</td>'
            f'<td class="text">{left}</td><td class="no">{new_no}</td>'
            f'<td class="text">{right}</td></tr>'
        )
    parts.append("</tbody></table>")
    parts.append("</main></body></html>")
    return "\n".join(parts) + "\n"

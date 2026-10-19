"""Table extraction from PDFs with a text layer (PyMuPDF find_tables)."""

from __future__ import annotations

import contextlib
import csv
import io
import os
from collections.abc import Callable
from dataclasses import dataclass, field
from pathlib import Path

import pymupdf
from openpyxl import Workbook
from openpyxl.comments import Comment

from jv_common import UserError, safe_output_path

from .numbers import parse_number

if hasattr(pymupdf, "no_recommend_layout"):
    pymupdf.no_recommend_layout()

NUMBER_FORMAT_DECIMALS = "#,##0.00"
NUMBER_FORMAT_INTEGER = "0"
COMMENT_AUTHOR = "jv pdf-tables"


@dataclass
class Table:
    page: int  # 1-based
    index: int  # 1-based on its page
    header: list[str]
    rows: list[list[str]]
    pages: list[int] = field(default_factory=list)

    @property
    def name(self) -> str:
        return f"s{self.page}_t{self.index}"

    @property
    def col_count(self) -> int:
        return len(self.header)


@dataclass
class Result:
    xlsx: Path | None
    csv_files: list[Path]
    tables: list[Table]
    page_count: int
    pages_scanned: list[int]
    warnings: list[str]


def parse_pages(spec: str | None, page_count: int) -> list[int]:
    """Parse "3-7", "2" or "1,3-5" into a sorted list of 1-based pages."""
    if not spec:
        return list(range(1, page_count + 1))
    pages: set[int] = set()
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        try:
            if "-" in part:
                a, b = part.split("-", 1)
                start, end = int(a), int(b)
            else:
                start = end = int(part)
        except ValueError:
            raise UserError(
                f"Ugyldig sidevalg «{spec}». Bruk for eksempel --pages 3-7 eller --pages 1,4."
            ) from None
        if start < 1 or end < start:
            raise UserError(f"Ugyldig sideområde «{part}» i --pages.")
        if end > page_count:
            raise UserError(
                f"Side {end} finnes ikke. Dokumentet har {page_count} sider."
            )
        pages.update(range(start, end + 1))
    if not pages:
        raise UserError(f"Ugyldig sidevalg «{spec}».")
    return sorted(pages)


def _clean(cell: str | None) -> str:
    if cell is None:
        return ""
    return " ".join(cell.split())


def _open_pdf(path: Path) -> pymupdf.Document:
    if not path.exists():
        raise UserError(f"Finner ikke filen «{path}».")
    if not path.is_file():
        raise UserError(f"«{path}» er ikke en fil.")
    if path.stat().st_size == 0:
        raise UserError(f"Filen «{path.name}» er tom.")
    try:
        doc = pymupdf.open(path, filetype="pdf")
    except Exception:  # noqa: BLE001 - any parse failure is a user error
        raise UserError(
            f"«{path.name}» kan ikke leses som PDF. Er det riktig filtype?"
        ) from None
    if doc.needs_pass:
        doc.close()
        raise UserError(f"«{path.name}» er passordbeskyttet.")
    if doc.page_count == 0:
        doc.close()
        raise UserError(f"«{path.name}» har ingen sider.")
    return doc


def extract_tables(
    path: str | os.PathLike[str],
    pages: str | None = None,
    progress: Callable[[float], None] | None = None,
) -> tuple[list[Table], int, list[int], list[str]]:
    """Find tables on the selected pages.

    Returns (tables, page_count, pages_scanned, warnings).
    """
    pdf = Path(path)
    doc = _open_pdf(pdf)
    warnings: list[str] = []
    tables: list[Table] = []
    try:
        selected = parse_pages(pages, doc.page_count)
        pages_without_text: list[int] = []
        for n, pno in enumerate(selected, start=1):
            page = doc[pno - 1]
            if not page.get_text("text").strip():
                pages_without_text.append(pno)
            else:
                # PyMuPDF may print hints to stdout; keep stdout clean for JSON.
                with contextlib.redirect_stdout(io.StringIO()):
                    found = page.find_tables().tables
                for idx, tab in enumerate(found, start=1):
                    raw = tab.extract()
                    if not raw:
                        continue
                    header = [_clean(c) for c in raw[0]]
                    rows = [[_clean(c) for c in r] for r in raw[1:]]
                    t = Table(page=pno, index=idx, header=header, rows=rows, pages=[pno])
                    tables.append(t)
                    none_counts = {sum(c is None for c in r) for r in raw}
                    if len(none_counts) > 1:
                        warnings.append(
                            f"Tabell {t.name} har ujevnt antall kolonner "
                            "(sammenslåtte eller manglende celler). Sjekk arket."
                        )
            if progress:
                progress(n / len(selected))
        if pages_without_text:
            if len(pages_without_text) == len(selected):
                warnings.append(
                    f"«{pdf.name}» mangler tekstlag. Kjør «jv ocr» først, "
                    "og bruk så pdf-tables på den nye filen."
                )
            else:
                listed = ", ".join(str(p) for p in pages_without_text)
                warnings.append(
                    f"Side {listed} mangler tekstlag og ble hoppet over. "
                    "Kjør «jv ocr» først for å få med disse."
                )
        return tables, doc.page_count, selected, warnings
    finally:
        doc.close()


def merge_tables(tables: list[Table]) -> list[Table]:
    """Join a table that ends a page with the first table on the next page
    when they have the same columns.

    Same columns means the same header (a repeated header row is dropped),
    or the same column count where the next part has no header and starts
    directly with numbers (that row is kept as data).
    """
    last_index = {}
    for t in tables:
        last_index[t.page] = max(last_index.get(t.page, 0), t.index)

    merged: list[Table] = []
    last_part: tuple[int, int] | None = None  # (page, index) of prev's last part
    for t in tables:
        prev = merged[-1] if merged else None
        if (
            prev is not None
            and last_part is not None
            and t.index == 1
            and t.page == last_part[0] + 1
            and last_part[1] == last_index[last_part[0]]
            and t.col_count == prev.col_count
        ):
            if t.header == prev.header:
                prev.rows.extend(t.rows)
                prev.pages.append(t.page)
                last_part = (t.page, t.index)
                continue
            if any(parse_number(h) is not None for h in t.header):
                prev.rows.append(list(t.header))
                prev.rows.extend(t.rows)
                prev.pages.append(t.page)
                last_part = (t.page, t.index)
                continue
        merged.append(
            Table(t.page, t.index, list(t.header), [list(r) for r in t.rows], list(t.pages))
        )
        last_part = (t.page, t.index)
    return merged


def _cell_value(text: str) -> tuple[object, str | None, bool]:
    """Return (value, comment_text, is_decimal)."""
    parsed = parse_number(text)
    if parsed is None:
        return (text if text else None), None, False
    comment = None
    if parsed.uncertain:
        comment = f"Usikker tolkning av tall. Originaltekst: «{text}»"
    return parsed.value, comment, isinstance(parsed.value, float) or "," in text


def write_xlsx(tables: list[Table], target: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for t in tables:
        ws = wb.create_sheet(title=t.name[:31])
        ws.append(t.header)
        for r_idx, row in enumerate(t.rows, start=2):
            for c_idx, text in enumerate(row, start=1):
                value, comment, is_decimal = _cell_value(text)
                cell = ws.cell(row=r_idx, column=c_idx, value=value)
                if isinstance(value, (int, float)):
                    cell.number_format = (
                        NUMBER_FORMAT_DECIMALS if is_decimal else NUMBER_FORMAT_INTEGER
                    )
                if comment:
                    cell.comment = Comment(comment, COMMENT_AUTHOR)
        for c_idx, _ in enumerate(t.header, start=1):
            width = max(
                [len(t.header[c_idx - 1])]
                + [len(r[c_idx - 1]) for r in t.rows if c_idx - 1 < len(r)]
            )
            ws.column_dimensions[ws.cell(row=1, column=c_idx).column_letter].width = min(
                60, width + 2
            )
    wb.save(target)


def write_csv(table: Table, target: Path) -> None:
    """CSV with ';' as separator (opens directly in Norwegian Excel), numbers
    as plain values with '.' as decimal sign."""
    with open(target, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(table.header)
        for row in table.rows:
            out = []
            for text in row:
                parsed = parse_number(text)
                out.append(text if parsed is None else repr(parsed.value))
            writer.writerow(out)


def run(
    path: str | os.PathLike[str],
    out_dir: str | os.PathLike[str],
    pages: str | None = None,
    merge: bool = False,
    write_csv_files: bool = False,
    progress: Callable[[float], None] | None = None,
) -> Result:
    tables, page_count, selected, warnings = extract_tables(path, pages, progress)
    if merge:
        tables = merge_tables(tables)
    if not tables:
        if not any("tekstlag" in w for w in warnings):
            warnings.append("Fant ingen tabeller på de valgte sidene.")
        return Result(None, [], [], page_count, selected, warnings)

    xlsx = safe_output_path(path, out_dir, "-tabeller.xlsx")
    written: list[Path] = []
    csv_files: list[Path] = []
    try:
        written.append(xlsx)
        write_xlsx(tables, xlsx)
        if write_csv_files:
            for t in tables:
                target = safe_output_path(path, out_dir, f"-{t.name}.csv")
                written.append(target)
                write_csv(t, target)
                csv_files.append(target)
    except BaseException:
        # never leave output files behind that no report mentions
        for f in written:
            with contextlib.suppress(OSError):
                f.unlink()
        raise
    return Result(xlsx, csv_files, tables, page_count, selected, warnings)


def validate_inputs(
    paths: list[str],
    out_dir: str | os.PathLike[str],
    pages: str | None = None,
) -> None:
    """Check every input before anything is written. Raises UserError."""
    targets: dict[Path, str] = {}
    for raw in paths:
        path = Path(raw)
        doc = _open_pdf(path)
        try:
            parse_pages(pages, doc.page_count)
        finally:
            doc.close()
        target = safe_output_path(path, out_dir, "-tabeller.xlsx")
        key = target.resolve()
        if key in targets:
            raise UserError(
                f"«{path.name}» og «{targets[key]}» ville gitt samme utfil. "
                "Kjør dem hver for seg eller gi dem ulike filnavn."
            )
        targets[key] = path.name

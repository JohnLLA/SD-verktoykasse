from openpyxl import load_workbook

from jv_pdf_tables.core import Table, merge_tables, write_xlsx


def test_uncertain_number_gets_comment(tmp_path):
    t = Table(page=1, index=1, header=["A", "B"], rows=[["x", "1.234.567,50"], ["y", "12"]], pages=[1])
    target = tmp_path / "t.xlsx"
    write_xlsx([t], target)
    ws = load_workbook(target)["s1_t1"]
    assert ws["B2"].value == 1234567.5
    assert "1.234.567,50" in ws["B2"].comment.text
    assert ws["B3"].value == 12 and ws["B3"].comment is None


def test_merge_without_repeated_header():
    a = Table(1, 1, ["Post", "Beløp"], [["a", "1"]], [1])
    b = Table(2, 1, ["b", "2"], [["c", "3"]], [2])
    merged = merge_tables([a, b])
    assert len(merged) == 1
    assert merged[0].rows == [["a", "1"], ["b", "2"], ["c", "3"]]
    assert merged[0].pages == [1, 2]


def test_no_merge_for_different_columns_or_not_last_table():
    a = Table(1, 1, ["A", "B"], [["a", "1"]], [1])
    b = Table(1, 2, ["X", "Y", "Z"], [["1", "2", "3"]], [1])
    c = Table(2, 1, ["A", "B"], [["c", "3"]], [2])
    merged = merge_tables([a, b, c])
    assert [t.name for t in merged] == ["s1_t1", "s1_t2", "s2_t1"]

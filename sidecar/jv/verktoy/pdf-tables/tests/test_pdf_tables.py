import csv
import hashlib
import json
import shutil
from pathlib import Path

import pytest
from openpyxl import load_workbook

from jv_pdf_tables.cli import main

DATA = Path(__file__).resolve().parents[3] / "testdata" / "pdf-tables"
BUDSJETT = DATA / "budsjett.pdf"
SKANNET = DATA / "skannet_tabell.pdf"
FASIT = json.loads((DATA / "budsjett_fasit.json").read_text(encoding="utf-8"))


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def run_tool(capsys, *args):
    code = main([str(a) for a in args])
    captured = capsys.readouterr()
    report = json.loads(captured.out) if code == 0 else None
    return code, report, captured


def xlsx_of(report):
    paths = [o["path"] for o in report["outputs"] if o["type"] == "xlsx"]
    assert len(paths) == 1
    return load_workbook(paths[0])


def check_table1(ws):
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0] == ("Post", "Regnskap 2025", "Budsjett 2026", "Endring")
    data = rows[1:]
    assert len(data) == len(FASIT["table1"])
    for row, want in zip(data, FASIT["table1"]):
        assert row[0] == want["post"]
        for got, key in zip(row[1:], ("regnskap_2025", "budsjett_2026", "endring")):
            assert isinstance(got, (int, float)), f"{want['post']} {key}: {got!r}"
            assert got == pytest.approx(want[key], abs=1e-9)


def check_table2_rows(data, fasit_rows):
    assert len(data) == len(fasit_rows)
    for row, want in zip(data, fasit_rows):
        assert str(row[0]) == want["ansvar"]
        assert row[1] == want["rammeomrade"]
        assert all(isinstance(v, (int, float)) for v in row[2:])
        assert row[2] == pytest.approx(want["budsjett_2026"], abs=1e-9)
        assert row[3] == pytest.approx(want["endring"], abs=1e-9)


def test_merge_gives_correct_values(tmp_path, capsys):
    code, report, _ = run_tool(capsys, BUDSJETT, "--out", tmp_path, "--merge")
    assert code == 0
    assert report["tool"] == "pdf-tables"
    assert report["inputs"][0]["sha256"] == sha(BUDSJETT)
    wb = xlsx_of(report)
    assert wb.sheetnames == ["s1_t1", "s1_t2"]
    check_table1(wb["s1_t1"])

    rows = list(wb["s1_t2"].iter_rows(values_only=True))
    assert rows[0] == ("Ansvar", "Rammeområde", "Budsjett 2026", "Endring fra 2025")
    data = rows[1:]
    assert len(data) == 50
    check_table2_rows(data, FASIT["table2"])
    assert sum(r[2] for r in data) == pytest.approx(57976226.25, abs=1e-6)
    assert sum(r[3] for r in data) == pytest.approx(-56787.00, abs=1e-6)

    sheet = next(s for s in report["outputs"][0]["sheets"] if s["name"] == "s1_t2")
    assert sheet["pages"] == FASIT["table2_pages"]
    assert report["warnings"] == []


def test_without_merge_one_sheet_per_table(tmp_path, capsys):
    code, report, _ = run_tool(capsys, BUDSJETT, "--out", tmp_path)
    assert code == 0
    wb = xlsx_of(report)
    assert wb.sheetnames == ["s1_t1", "s1_t2", "s2_t1"]
    check_table1(wb["s1_t1"])
    part1 = list(wb["s1_t2"].iter_rows(values_only=True))[1:]
    part2 = list(wb["s2_t1"].iter_rows(values_only=True))
    assert part2[0] == ("Ansvar", "Rammeområde", "Budsjett 2026", "Endring fra 2025")
    assert len(part1) == 24 and len(part2[1:]) == 26
    check_table2_rows(part1 + part2[1:], FASIT["table2"])


def test_no_comments_for_certain_numbers(tmp_path, capsys):
    _, report, _ = run_tool(capsys, BUDSJETT, "--out", tmp_path, "--merge")
    wb = xlsx_of(report)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                assert cell.comment is None


def test_pages_option(tmp_path, capsys):
    code, report, _ = run_tool(capsys, BUDSJETT, "--out", tmp_path, "--pages", "2")
    assert code == 0
    wb = xlsx_of(report)
    assert wb.sheetnames == ["s2_t1"]
    data = list(wb["s2_t1"].iter_rows(values_only=True))[1:]
    check_table2_rows(data, FASIT["table2"][24:])


@pytest.mark.parametrize("pages", ["3", "0-1", "2-1", "x"])
def test_bad_pages(tmp_path, capsys, pages):
    code, _, captured = run_tool(capsys, BUDSJETT, "--out", tmp_path, "--pages", pages)
    assert code == 1
    assert captured.err.startswith("Feil:")
    assert "Traceback" not in captured.err


def test_csv(tmp_path, capsys):
    code, report, _ = run_tool(capsys, BUDSJETT, "--out", tmp_path, "--merge", "--csv")
    assert code == 0
    csvs = [Path(o["path"]) for o in report["outputs"] if o["type"] == "csv"]
    assert [p.name for p in csvs] == ["budsjett-s1_t1.csv", "budsjett-s1_t2.csv"]
    with open(csvs[1], encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.reader(handle, delimiter=";"))
    assert len(rows) == 51
    assert float(rows[1][3]) == pytest.approx(-24690.0)
    assert sum(float(r[2]) for r in rows[1:]) == pytest.approx(57976226.25, abs=1e-6)


def test_scanned_pdf_warns_about_ocr(tmp_path, capsys):
    code, report, _ = run_tool(capsys, SKANNET, "--out", tmp_path)
    assert code == 0
    assert report["outputs"] == []
    assert any("ocr" in w and "tekstlag" in w for w in report["warnings"])
    assert list(tmp_path.iterdir()) == []


def test_unicode_and_spaces_in_filename(tmp_path, capsys):
    src_dir = tmp_path / "inn mappe"
    src_dir.mkdir()
    src = src_dir / "Budsjett Ærøy; 2026.pdf"
    shutil.copy(BUDSJETT, src)
    before = sha(src)
    out = tmp_path / "ut æøå"
    code, report, _ = run_tool(capsys, src, "--out", out, "--merge")
    assert code == 0
    assert Path(report["outputs"][0]["path"]).name == "Budsjett Ærøy; 2026-tabeller.xlsx"
    assert sha(src) == before


def test_input_unchanged_and_no_overwrite(tmp_path, capsys):
    src = tmp_path / "budsjett.pdf"
    shutil.copy(BUDSJETT, src)
    before = sha(src)
    _, r1, _ = run_tool(capsys, src, "--out", tmp_path)
    _, r2, _ = run_tool(capsys, src, "--out", tmp_path)
    assert sha(src) == before
    assert r1["outputs"][0]["path"] != r2["outputs"][0]["path"]
    assert Path(r2["outputs"][0]["path"]).name == "budsjett-tabeller-2.xlsx"


def test_missing_file(tmp_path, capsys):
    code, _, captured = run_tool(capsys, tmp_path / "finnes_ikke.pdf", "--out", tmp_path)
    assert code == 1
    assert "Finner ikke filen" in captured.err
    assert "Traceback" not in captured.err


def test_empty_file(tmp_path, capsys):
    empty = tmp_path / "tom.pdf"
    empty.write_bytes(b"")
    code, _, captured = run_tool(capsys, empty, "--out", tmp_path / "ut")
    assert code == 1
    assert "tom" in captured.err
    assert "Traceback" not in captured.err


def test_wrong_format(tmp_path, capsys):
    fake = tmp_path / "notat.pdf"
    fake.write_text("Dette er ikke en PDF.", encoding="utf-8")
    code, _, captured = run_tool(capsys, fake, "--out", tmp_path / "ut")
    assert code == 1
    assert "kan ikke leses som PDF" in captured.err
    assert "Traceback" not in captured.err


def test_missing_out_argument(capsys):
    code, _, captured = run_tool(capsys, BUDSJETT)
    assert code == 1
    assert captured.err.startswith("Feil:")


def test_stdout_is_only_json(tmp_path, capsys):
    code, _, captured = run_tool(capsys, BUDSJETT, "--out", tmp_path)
    assert code == 0
    json.loads(captured.out)
    progress = [json.loads(line) for line in captured.err.splitlines() if line.strip()]
    assert progress and all("progress" in p for p in progress)


def _uneven_pdf(path):
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    x = [72, 200, 300, 400]
    y = [100, 120, 140, 160]
    cells = [
        [(0, 1), (1, 2), (2, 3)],
        [(0, 1), (1, 3)],  # merged cell over two columns
        [(0, 1), (1, 2), (2, 3)],
    ]
    texts = [["Post", "A", "B"], ["Sum", "1 000,00"], ["Rest", "2", "3"]]
    for r, row in enumerate(cells):
        for (c0, c1), text in zip(row, texts[r]):
            rect = pymupdf.Rect(x[c0], y[r], x[c1], y[r + 1])
            page.draw_rect(rect, color=(0, 0, 0), width=0.8)
            page.insert_text((rect.x0 + 3, rect.y1 - 5), text, fontsize=9)
    doc.save(path)
    doc.close()


def test_uneven_columns_warning(tmp_path, capsys):
    src = tmp_path / "ujevn.pdf"
    _uneven_pdf(src)
    code, report, _ = run_tool(capsys, src, "--out", tmp_path / "ut")
    assert code == 0
    assert any("ujevnt antall kolonner" in w for w in report["warnings"])


def test_alle_filer_valideres_for_noe_skrives(tmp_path, capsys):
    out = tmp_path / "ut"
    code, _, cap = run_tool(capsys, BUDSJETT, tmp_path / "finnes_ikke.pdf", "--out", out)
    assert code == 1
    assert "Finner ikke filen" in cap.err
    assert not out.exists() or list(out.iterdir()) == []


def test_ugyldig_sidevalg_stopper_for_skriving(tmp_path, capsys):
    out = tmp_path / "ut"
    code, _, cap = run_tool(capsys, BUDSJETT, SKANNET, "--pages", "999", "--out", out)
    assert code == 1
    assert "finnes ikke" in cap.err
    assert not out.exists() or list(out.iterdir()) == []


def test_samme_utfil_to_ganger_avvises(tmp_path, capsys):
    a = tmp_path / "a"
    b = tmp_path / "b"
    a.mkdir()
    b.mkdir()
    shutil.copy(BUDSJETT, a / "x.pdf")
    shutil.copy(BUDSJETT, b / "x.pdf")
    out = tmp_path / "ut"
    code, _, cap = run_tool(capsys, a / "x.pdf", b / "x.pdf", "--out", out)
    assert code == 1
    assert "samme utfil" in cap.err
    assert not out.exists() or list(out.iterdir()) == []


def test_feil_midt_i_behandlingen_gir_rapport_med_fullforte(tmp_path, capsys, monkeypatch):
    import jv_pdf_tables.core as core

    second = tmp_path / "andre fil æøå.pdf"
    shutil.copy(BUDSJETT, second)
    real = core.write_xlsx
    calls = []

    def flaky(tables, target):
        calls.append(target)
        if len(calls) == 2:
            target.write_bytes(b"halv")
            raise PermissionError(13, "nei")
        real(tables, target)

    monkeypatch.setattr(core, "write_xlsx", flaky)
    out = tmp_path / "ut"
    code = main([str(BUDSJETT), str(second), "--out", str(out)])
    cap = capsys.readouterr()
    assert code == 1
    assert "Feil: Filsystemet nektet tilgang" in cap.err
    report = json.loads(cap.out)
    assert len(report["outputs"]) == 1
    assert [i["path"] for i in report["inputs"]] == [str(BUDSJETT)]
    assert any("stoppet" in w for w in report["warnings"])
    assert sorted(p.name for p in out.iterdir()) == [Path(report["outputs"][0]["path"]).name]

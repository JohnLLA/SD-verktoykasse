"""Generate all synthetic test files in testdata/ from scratch.

All persons, numbers and texts are fictitious. Run from the repo root:

    uv run --no-project --python 3.12 --with reportlab --with pillow \
        --with python-docx --with openpyxl --with pymupdf \
        python scripts/lag_testdata.py

Requires ffmpeg, ffprobe and exiftool in PATH. No network access is used.
"""

from __future__ import annotations

import json
import re
import shutil
import subprocess
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import pymupdf as fitz
from docx import Document
from docx.shared import Pt
from PIL import Image, ImageDraw, ImageFont
from reportlab import rl_config
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

rl_config.invariant = 1  # reproducible PDFs (fixed dates and IDs)

ROOT = Path(__file__).resolve().parent.parent
TD = ROOT / "testdata"
FIXED_DATE = datetime(2026, 1, 15, 10, 30, 0, tzinfo=timezone.utc)

# ---------------------------------------------------------------------------
# Fonts: use a local TTF (no downloads) so that æ, ø, å and U+2212 render.
# ---------------------------------------------------------------------------
FONT_CANDIDATES = [
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
    ("/System/Library/Fonts/Supplemental/Arial.ttf",
     "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
    ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf"),
]
FONT = "Helvetica"
FONT_BOLD = "Helvetica-Bold"
FONT_FILE = None
for reg, bold in FONT_CANDIDATES:
    if Path(reg).exists() and Path(bold).exists():
        pdfmetrics.registerFont(TTFont("TDSans", reg))
        pdfmetrics.registerFont(TTFont("TDSans-Bold", bold))
        FONT, FONT_BOLD, FONT_FILE = "TDSans", "TDSans-Bold", reg
        break
MINUS = "\u2212" if FONT_FILE else "-"


def run(cmd: list[str]) -> subprocess.CompletedProcess:
    res = subprocess.run(cmd, capture_output=True, text=True)
    if res.returncode != 0:
        sys.stderr.write(res.stderr[-3000:])
        raise RuntimeError(f"Kommando feilet: {' '.join(cmd[:4])} ...")
    return res


def styles():
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName=FONT,
                          fontSize=11, leading=15, spaceAfter=8)
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontName=FONT_BOLD,
                        fontSize=16, leading=20, spaceAfter=10)
    return body, h1


def text_pdf(path: Path, title: str, paragraphs: list[str], meta: dict | None = None):
    body, h1 = styles()
    meta = meta or {}
    doc = SimpleDocTemplate(str(path), pagesize=A4, title=meta.get("title", title),
                            author=meta.get("author", "Testdata"),
                            subject=meta.get("subject", ""),
                            creator=meta.get("creator", "lag_testdata.py"),
                            leftMargin=20 * mm, rightMargin=20 * mm)
    flow = [Paragraph(title, h1)] + [Paragraph(p, body) for p in paragraphs]
    doc.build(flow)


def pdf_text(path: Path) -> list[str]:
    with fitz.open(path) as d:
        return [p.get_text() for p in d]


def norm(s: str) -> str:
    return re.sub(r"\s+", " ", s).strip()


# ---------------------------------------------------------------------------
# 1. compare
# ---------------------------------------------------------------------------
VEDTAK_TITLE = "Vedtak i sak 42/26 – Tilskudd til Fiktivhavn kulturhus"
VEDTAK_V1 = [
    "Formannskapet i Fiktivhavn kommune behandlet saken i møte 12. februar 2026.",
    "Kommunen innvilger et tilskudd på 1 200 000 kroner til rehabilitering av "
    "kulturhuset. Tilskuddet utbetales i to like deler. Første del utbetales når "
    "avtalen er signert.",
    "Tilskuddet skal brukes til nytt ventilasjonsanlegg, universell utforming av "
    "inngangspartiet og oppgradering av scenelys.",
    "Stiftelsen Fiktivhavn kulturhus skal levere en sluttrapport med revidert "
    "regnskap innen 1. desember 2027.",
    "Ubrukte midler skal betales tilbake til kommunen. Kommunen kan kreve "
    "tilbakebetaling dersom vilkårene ikke er oppfylt.",
    "Vedtaket kan påklages til Statsforvalteren innen tre uker fra det er mottatt.",
]
# v2: sentence "Tilskuddet utbetales i to like deler." moved from paragraph 2 to
# the end of paragraph 4, amount 1 200 000 -> 1 250 000, paragraph 5 deleted.
MOVED_SENTENCE = "Tilskuddet utbetales i to like deler."
VEDTAK_V2 = [
    VEDTAK_V1[0],
    "Kommunen innvilger et tilskudd på 1 250 000 kroner til rehabilitering av "
    "kulturhuset. Første del utbetales når avtalen er signert.",
    VEDTAK_V1[2],
    VEDTAK_V1[3] + " " + MOVED_SENTENCE,
    VEDTAK_V1[5],
]


def make_compare():
    out = TD / "compare"
    out.mkdir(parents=True, exist_ok=True)
    for ver, paras in (("v1", VEDTAK_V1), ("v2", VEDTAK_V2)):
        d = Document()
        d.core_properties.author = "Testdata"
        d.core_properties.created = FIXED_DATE.replace(tzinfo=None)
        d.add_heading(VEDTAK_TITLE, level=1)
        for p in paras:
            d.add_paragraph(p)
        d.save(out / f"vedtak_{ver}.docx")
        (out / f"vedtak_{ver}.txt").write_text(
            VEDTAK_TITLE + "\n\n" + "\n\n".join(paras) + "\n", encoding="utf-8")
        text_pdf(out / f"vedtak_{ver}.pdf", VEDTAK_TITLE, paras)


# ---------------------------------------------------------------------------
# 2. metadata
# ---------------------------------------------------------------------------
BRYGGEN = (60.39745, 5.32374)  # Bryggen, Bergen (public landmark)


def load_font(size: int):
    if FONT_FILE:
        try:
            return ImageFont.truetype(FONT_FILE, size)
        except OSError:
            pass
    return ImageFont.load_default()


def make_metadata():
    out = TD / "metadata"
    out.mkdir(parents=True, exist_ok=True)
    # Stored landscape 640x480. With Orientation=6 the viewer rotates 90° CW,
    # so the stored LEFT edge becomes the displayed TOP. We draw an arrow and
    # the word TOPP pointing to the stored left edge, rotated accordingly.
    w, h = 640, 480
    img = Image.new("RGB", (w, h), (40, 90, 160))
    dr = ImageDraw.Draw(img)
    dr.rectangle([0, 0, 60, h], fill=(230, 60, 40))  # red band = displayed top
    dr.polygon([(80, h // 2), (200, h // 2 - 80), (200, h // 2 + 80)], fill="white")
    dr.rectangle([200, h // 2 - 25, 420, h // 2 + 25], fill="white")
    label = Image.new("RGBA", (360, 90), (0, 0, 0, 0))
    ImageDraw.Draw(label).text((10, 5), "TOPP", font=load_font(72), fill="yellow")
    label = label.rotate(90, expand=True)  # CCW in storage, upright after 90° CW display
    img.paste(label, (460, 60), label)
    path = out / "bilde_gps.jpg"
    img.save(path, "JPEG", quality=85)
    lat, lon = BRYGGEN
    run(["exiftool", "-overwrite_original", "-n",
         "-Make=Testkamera", "-Model=Fiktiv X100",
         "-Software=lag_testdata.py",
         "-Orientation=6",
         "-DateTimeOriginal=2026:01:15 10:30:00",
         f"-GPSLatitude={lat}", "-GPSLatitudeRef=N",
         f"-GPSLongitude={lon}", "-GPSLongitudeRef=E",
         "-GPSAltitude=4", "-GPSAltitudeRef=0",
         "-Artist=Ola Testesen",
         str(path)])

    # DOCX with author, last modified by and company (app.xml).
    paras = ["Dette er et oppdiktet notat laget for å teste metadata-verktøyet.",
             "Dokumentet har forfatter, firma og «sist endret av» i metadataene."]
    d = Document()
    cp = d.core_properties
    cp.author = "Ola Testesen"
    cp.last_modified_by = "Kari Fiktivsen"
    cp.title = "Internt notat"
    cp.comments = "Oppdiktet testdokument"
    cp.created = FIXED_DATE.replace(tzinfo=None)
    cp.modified = FIXED_DATE.replace(tzinfo=None)
    d.add_heading("Internt notat", level=1)
    for p in paras:
        d.add_paragraph(p)
    tmp = out / "_tmp.docx"
    d.save(tmp)
    app_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" '
        'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
        "<Template>Normal.dotm</Template><TotalTime>12</TotalTime>"
        "<Application>Microsoft Office Word</Application>"
        "<Company>Eksempelfirma AS</Company><Manager>Per Prøvesen</Manager>"
        "<AppVersion>16.0000</AppVersion></Properties>"
    )
    target = out / "dokument_forfatter.docx"
    with zipfile.ZipFile(tmp) as zin, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as zout:
        names = zin.namelist()
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "docProps/app.xml":
                data = app_xml.encode("utf-8")
            zout.writestr(item, data)
        if "docProps/app.xml" not in names:
            zout.writestr("docProps/app.xml", app_xml)
    tmp.unlink()

    # PDF with Info dictionary + XMP.
    pdf = out / "dokument_forfatter.pdf"
    tmp_pdf = out / "_tmp.pdf"
    text_pdf(tmp_pdf, "Internt notat", paras,
             meta={"title": "Internt notat", "author": "Ola Testesen",
                   "subject": "Oppdiktet testdokument", "creator": "Microsoft Word"})
    with fitz.open(tmp_pdf) as doc:
        doc.set_metadata({
            "title": "Internt notat", "author": "Ola Testesen",
            "subject": "Oppdiktet testdokument", "keywords": "test, metadata",
            "creator": "Microsoft Word", "producer": "lag_testdata.py",
            "creationDate": "D:20260115103000Z", "modDate": "D:20260115103000Z",
        })
        info = doc.xref_get_key(-1, "Info")[1]
        info_xref = int(info.split()[0])
        doc.xref_set_key(info_xref, "Company", fitz.get_pdf_str("Eksempelfirma AS"))
        doc.xref_set_key(info_xref, "LastModifiedBy", fitz.get_pdf_str("Kari Fiktivsen"))
        xmp = (
            '<?xpacket begin="\ufeff" id="W5M0MpCehiHzreSzNTczkc9d"?>'
            '<x:xmpmeta xmlns:x="adobe:ns:meta/"><rdf:RDF '
            'xmlns:rdf="http://www.w3.org/1999/02/22-rdf-syntax-ns#">'
            '<rdf:Description rdf:about="" xmlns:dc="http://purl.org/dc/elements/1.1/" '
            'xmlns:pdf="http://ns.adobe.com/pdf/1.3/" '
            'xmlns:xmp="http://ns.adobe.com/xap/1.0/">'
            "<dc:creator><rdf:Seq><rdf:li>Ola Testesen</rdf:li></rdf:Seq></dc:creator>"
            '<dc:title><rdf:Alt><rdf:li xml:lang="x-default">Internt notat</rdf:li></rdf:Alt></dc:title>'
            "<pdf:Producer>lag_testdata.py</pdf:Producer>"
            "<xmp:CreatorTool>Microsoft Word</xmp:CreatorTool>"
            "</rdf:Description></rdf:RDF></x:xmpmeta>"
            '<?xpacket end="w"?>'
        )
        doc.set_xml_metadata(xmp)
        doc.save(pdf, garbage=3, deflate=True)
    tmp_pdf.unlink()


# ---------------------------------------------------------------------------
# 3. keyframes
# ---------------------------------------------------------------------------
def make_keyframes():
    out = TD / "keyframes"
    out.mkdir(parents=True, exist_ok=True)
    # Hard cuts at 2, 5 and 8 s: black 0-2, white 2-5, blue 5-8, yellow 8-10.
    segs = [("black", 2), ("white", 3), ("blue", 3), ("yellow", 2)]
    inputs, chains = [], []
    for i, (c, dur) in enumerate(segs):
        inputs += ["-f", "lavfi", "-i", f"color=c={c}:s=320x240:r=25:d={dur}"]
        chains.append(f"[{i}:v]")
    fc = "".join(chains) + f"concat=n={len(segs)}:v=1:a=0,format=yuv420p[v]"
    run(["ffmpeg", "-y", "-hide_banner", "-loglevel", "error", *inputs,
         "-filter_complex", fc, "-map", "[v]", "-c:v", "libx264",
         "-preset", "veryfast", "-g", "25", "-crf", "28",
         "-movflags", "+faststart", str(out / "klipp.mp4")])
    # Portrait phone video: stored 320x180 landscape, rotation metadata 90°.
    tmp = out / "_tmp.mov"
    run(["ffmpeg", "-y", "-hide_banner", "-loglevel", "error",
         "-f", "lavfi", "-i", "testsrc=s=320x180:r=25:d=4",
         "-vf", "format=yuv420p", "-c:v", "libx264", "-preset", "veryfast",
         "-crf", "30", str(tmp)])
    # Goal (iPhone convention): ffprobe shows tag rotate=90 and a display
    # matrix with rotation -90, i.e. display 90 degrees clockwise. Some ffmpeg
    # builds (e.g. 4.4 here) invert the sign when muxing the tag, so we check
    # the result and fall back to the other value.
    dst = out / "stående.mov"
    for value in ("90", "270"):
        run(["ffmpeg", "-y", "-hide_banner", "-loglevel", "error", "-i", str(tmp),
             "-c", "copy", "-metadata:s:v:0", f"rotate={value}", str(dst)])
        st = json.loads(run(["ffprobe", "-v", "error", "-of", "json", "-show_streams",
                             "-select_streams", "v:0", str(dst)]).stdout)["streams"][0]
        if st.get("tags", {}).get("rotate") == "90":
            break
    tmp.unlink()


# ---------------------------------------------------------------------------
# 4. loudness
# ---------------------------------------------------------------------------
def measure(path: Path) -> tuple[float, float]:
    res = subprocess.run(["ffmpeg", "-hide_banner", "-nostats", "-i", str(path),
                          "-af", "ebur128=peak=true", "-f", "null", "-"],
                         capture_output=True, text=True)
    tail = res.stderr[res.stderr.rfind("Summary:"):]
    i = float(re.search(r"I:\s+(-?[\d.]+) LUFS", tail).group(1))
    tp = float(re.search(r"Peak:\s+(-?[\d.]+) dBFS", tail).group(1))
    return i, tp


SPEECHLIKE = ("anoisesrc=color=pink:seed=42:duration={d}:sample_rate=48000,"
              "highpass=f=200,lowpass=f=4000,tremolo=f=4:d=0.6")


def render_level(path: Path, target: float, limit: bool, extra: list[str] | None = None):
    gain = target + 21.0  # source is about -21 LUFS at 0 dB gain
    for _ in range(12):
        af = SPEECHLIKE.format(d=20) + f",volume={gain:.2f}dB"
        if limit:
            af += ",alimiter=limit=0.89:level=0:attack=1:release=50"
        run(["ffmpeg", "-y", "-hide_banner", "-loglevel", "error",
             "-f", "lavfi", "-i", af, "-ac", "1", "-c:a", "pcm_s16le",
             *(extra or []), str(path)])
        i, tp = measure(path)
        if abs(i - target) <= 0.2:
            break
        gain += (target - i) * (2.0 if limit else 1.0)
    return measure(path)


LOUD = {}


def make_loudness():
    out = TD / "loudness"
    out.mkdir(parents=True, exist_ok=True)
    LOUD["lav.wav"] = render_level(out / "lav.wav", -30.0, limit=False)
    LOUD["høy.wav"] = render_level(out / "høy.wav", -8.0, limit=True)
    # Video with audio track (about -20 LUFS), 10 s.
    run(["ffmpeg", "-y", "-hide_banner", "-loglevel", "error",
         "-f", "lavfi", "-i", "testsrc=s=320x240:r=25:d=10",
         "-f", "lavfi", "-i", SPEECHLIKE.format(d=10) + ",volume=1dB",
         "-c:v", "libx264", "-preset", "veryfast", "-crf", "30", "-pix_fmt", "yuv420p",
         "-c:a", "aac", "-b:a", "96k", "-ar", "48000", "-ac", "2", "-shortest",
         str(out / "video_med_lyd.mp4")])
    LOUD["video_med_lyd.mp4"] = measure(out / "video_med_lyd.mp4")


# ---------------------------------------------------------------------------
# 5. subtitles
# ---------------------------------------------------------------------------
TRANSCRIPT = [
    "Velkommen til Lokalpodden, der vi snakker om det som skjer i Fiktivhavn.",
    "I dag har vi med oss kulturhussjef Mette Oppdiktsen.",
    "Hei, og velkommen.",
    "Takk.",
    "Kulturhuset har fått et tilskudd fra kommunen, og det har vært mye diskusjon "
    "om pengene i formannskapet, i lokalavisen og i sosiale medier de siste ukene, "
    "så la oss begynne der.",
    "Hva skal pengene brukes til?",
    "Først og fremst til ventilasjon.",
    "Det har vært et problem i mange år, særlig om vinteren når salen er full.",
    "Folk har klaget på at det blir varmt og tett.",
    "Ja.",
    "Og så skal vi gjøre inngangen tilgjengelig for alle, med ny rampe og "
    "automatiske dører.",
    "Hvor mye koster det?",
    "Totalt regner vi med litt over to millioner kroner.",
    "Kommunen dekker en god del, og resten kommer fra stiftelsen og fra et fond.",
    "Er du fornøyd med beløpet?",
    "Nei, ikke helt.",
    "Vi søkte om mer, men vi forstår at kommunen har en stram økonomi.",
    "Så hva gjør dere hvis det ikke strekker til?",
    "Da må vi prioritere.",
    "Ventilasjonen kommer først, fordi det handler om helse og arbeidsmiljø.",
    "Scenelyset kan vente et år om nødvendig.",
    "Det er noen som mener at kulturhuset burde ha vært lagt ned i stedet.",
    "Hva tenker du om det?",
    "Jeg skjønner at folk stiller spørsmål.",
    "Men kulturhuset er det eneste stedet i bygda der korpset, teatergruppa, "
    "barnehagen og pensjonistforeningen kan være under samme tak.",
    "Hvis vi legger det ned, mister vi en møteplass.",
    "Mm.",
    "Hvor mange bruker huset i løpet av et år?",
    "Vi teller rundt tolv tusen besøk, og det er mye i en kommune med fire tusen "
    "innbyggere.",
    "Det er imponerende.",
    "Ja, vi er stolte av det.",
    "Når starter arbeidet?",
    "Vi håper å starte i april, etter påske.",
    "Da må vi stenge salen i omtrent tre måneder.",
    "Hva skjer med arrangementene i mellomtiden?",
    "Noen flytter vi til skolen, og noen utsetter vi til høsten.",
    "Korpset har allerede sagt at de kan øve i gymsalen.",
    "Til slutt, hva vil du si til dem som er skeptiske?",
    "Kom og se huset.",
    "Gå en runde med meg, så skal jeg vise dere hvorfor dette er nødvendig.",
    "Døren står åpen.",
    "Takk for at du kom.",
    "Takk selv.",
    "Det var alt for i dag.",
    "Neste uke snakker vi med ordføreren om den nye kaia og om hvorfor fergen "
    "fortsatt er forsinket hver fredag.",
    "Ha en fin uke.",
]
TARGET_DURATION = 180.0


def make_subtitles():
    out = TD / "subtitles"
    out.mkdir(parents=True, exist_ok=True)
    gap = 0.4
    raw = [max(0.7, len(t) / 13.0) for t in TRANSCRIPT]
    scale = (TARGET_DURATION - gap * (len(raw) - 1) - 0.5) / sum(raw)
    t = 0.5
    segments = []
    for idx, (text, dur) in enumerate(zip(TRANSCRIPT, raw)):
        dur *= scale
        start, end = round(t, 2), round(t + dur, 2)
        words = []
        toks = text.split(" ")
        weights = [len(w) + 1 for w in toks]
        total = sum(weights)
        wt = start
        for i, (w, wgt) in enumerate(zip(toks, weights)):
            wd = (end - start) * wgt / total
            ws, we = round(wt, 2), round(wt + wd, 2) if i < len(toks) - 1 else end
            words.append({"word": " " + w, "start": ws, "end": we,
                          "probability": round(0.85 + 0.1 * ((i * 7) % 10) / 10, 3)})
            wt += wd
        segments.append({
            "id": idx, "seek": int(start * 100) // 3000 * 3000, "start": start,
            "end": end, "text": " " + text, "tokens": [],
            "temperature": 0.0, "avg_logprob": -0.25, "compression_ratio": 1.4,
            "no_speech_prob": 0.01, "words": words,
        })
        t = end + gap
    data = {"text": " ".join(" " + s for s in TRANSCRIPT).strip(),
            "segments": segments, "language": "no"}
    (out / "transkripsjon.json").write_text(
        json.dumps(data, ensure_ascii=False, indent=1), encoding="utf-8")
    no_words = json.loads(json.dumps(data))
    for s in no_words["segments"]:
        del s["words"]
    (out / "transkripsjon_uten_ord.json").write_text(
        json.dumps(no_words, ensure_ascii=False, indent=1), encoding="utf-8")
    return data


# ---------------------------------------------------------------------------
# 6. pdf-tables
# ---------------------------------------------------------------------------
def nok(v: float, style: str = "minus") -> str:
    """Format as Norwegian number: 1 234 567,50. Negative as (x), −x or -x."""
    s = f"{abs(v):,.2f}".replace(",", " ").replace(".", ",")
    if v < 0:
        if style == "paren":
            return f"({s})"
        if style == "hyphen":
            return f"-{s}"
        return f"{MINUS}{s}"
    return s


TABLE1 = [  # post, regnskap 2025, budsjett 2026, endring (negative style)
    ("Skatt på inntekt og formue", 1234567.50, 1301000.00, "minus"),
    ("Rammetilskudd", 2450000.00, 2398500.25, "paren"),
    ("Brukerbetalinger", 312400.75, 298000.00, "paren"),
    ("Salgsinntekter", 87650.00, 91200.00, "minus"),
    ("Renteutgifter", -145320.40, -162000.00, "minus"),
    ("Avdrag på lån", -410000.00, -398250.60, "hyphen"),
]
RAMMER = ["Oppvekst", "Helse og omsorg", "Kultur", "Teknisk drift", "Sentraladministrasjon",
          "Næring", "Barnevern", "NAV", "Brann og redning", "Havn"]


def make_pdf_tables():
    out = TD / "pdf-tables"
    out.mkdir(parents=True, exist_ok=True)
    body, h1 = styles()
    cell = ParagraphStyle("cell", parent=body, fontSize=9, leading=11, spaceAfter=0)
    expected = {"table1": [], "table2": []}

    hdr = ["Post", "Regnskap 2025", "Budsjett 2026", "Endring"]
    rows1 = [hdr]
    for post, a, b, sty in TABLE1:
        diff = round(b - a, 2)
        rows1.append([post, nok(a, sty), nok(b, sty), nok(diff, sty)])
        expected["table1"].append({"post": post, "regnskap_2025": a,
                                   "budsjett_2026": b, "endring": diff,
                                   "tekst": rows1[-1][1:]})
    tstyle = TableStyle([
        ("FONTNAME", (0, 0), (-1, -1), FONT),
        ("FONTNAME", (0, 0), (-1, 0), FONT_BOLD),
        ("FONTSIZE", (0, 0), (-1, -1), 9),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.black),
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
    ])
    t1 = Table(rows1, colWidths=[60 * mm, 35 * mm, 35 * mm, 35 * mm])
    t1.setStyle(tstyle)

    hdr2 = ["Ansvar", "Rammeområde", "Budsjett 2026", "Endring fra 2025"]
    rows2 = [hdr2]
    for i in range(50):
        area = RAMMER[i % len(RAMMER)]
        ansvar = f"{1000 + i * 10}"
        budget = round(250000 + i * 37123.45, 2)
        change = round(((i * 7919) % 41 - 20) * 1234.5, 2)
        sty = "paren" if i % 3 == 0 else "minus"
        rows2.append([ansvar, area, nok(budget), nok(change, sty)])
        expected["table2"].append({"ansvar": ansvar, "rammeomrade": area,
                                   "budsjett_2026": budget, "endring": change,
                                   "tekst": rows2[-1][2:]})
    t2 = Table(rows2, colWidths=[25 * mm, 55 * mm, 40 * mm, 40 * mm], repeatRows=1)
    t2.setStyle(tstyle)

    doc = SimpleDocTemplate(str(out / "budsjett.pdf"), pagesize=A4,
                            title="Budsjett 2026 – Fiktivhavn kommune",
                            author="Testdata", leftMargin=20 * mm, rightMargin=20 * mm)
    flow = [Paragraph("Budsjett 2026 – Fiktivhavn kommune", h1),
            Paragraph("Tabell 1: Hovedoversikt drift (kroner). Negative tall vises "
                      "med minustegn eller i parentes.", body),
            t1, Spacer(1, 12),
            Paragraph("Tabell 2: Fordeling per ansvar (kroner). Tabellen fortsetter "
                      "på neste side.", body),
            t2]
    doc.build(flow)
    with fitz.open(out / "budsjett.pdf") as d:
        expected["pages"] = d.page_count
        pages_t2 = [i + 1 for i, p in enumerate(d) if "Ansvar" in p.get_text()]
    expected["table2_pages"] = pages_t2
    (out / "budsjett_fasit.json").write_text(
        json.dumps(expected, ensure_ascii=False, indent=1), encoding="utf-8")

    # Image-only version of page 1.
    image_only_pdf(out / "budsjett.pdf", out / "skannet_tabell.pdf", pages=[0], dpi=150)
    return expected


# ---------------------------------------------------------------------------
# Raster helpers (simulate scans)
# ---------------------------------------------------------------------------
def raster_page(src: Path, pno: int, dpi: int) -> Image.Image:
    with fitz.open(src) as d:
        pix = d[pno].get_pixmap(dpi=dpi, colorspace=fitz.csGRAY)
        return Image.frombytes("L", (pix.width, pix.height), pix.samples)


def images_to_pdf(images: list[Image.Image], dst: Path, dpi: int, text_pages: dict | None = None):
    """Build PDF; images become image-only pages; text_pages maps index->source page."""
    out = fitz.open()
    text_pages = text_pages or {}
    for i, img in enumerate(images):
        if i in text_pages:
            src, pno = text_pages[i]
            with fitz.open(src) as s:
                out.insert_pdf(s, from_page=pno, to_page=pno)
            continue
        w_pt, h_pt = img.width * 72 / dpi, img.height * 72 / dpi
        page = out.new_page(width=w_pt, height=h_pt)
        import io
        buf = io.BytesIO()
        img.save(buf, "PNG", optimize=True)
        page.insert_image(page.rect, stream=buf.getvalue())
    out.set_metadata({"producer": "lag_testdata.py (simulert skann)"})
    out.save(dst, garbage=3, deflate=True)
    out.close()


def image_only_pdf(src: Path, dst: Path, pages: list[int], dpi: int):
    images_to_pdf([raster_page(src, p, dpi) for p in pages], dst, dpi)


# ---------------------------------------------------------------------------
# 7. ocr
# ---------------------------------------------------------------------------
OCR_TITLE = "Bergen kommune – Melding om vedtak"
OCR_TEXT = [
    "Bergen kommune har mottatt søknaden din om tillatelse til å flytte "
    "sjøboden på Fiktivkaien. Søknaden er behandlet etter plan- og bygningsloven.",
    "Vedtak: Søknaden godkjennes. Arbeidet må være ferdig innen ett år. "
    "Tiltaket må ikke hindre ferdsel langs sjøen.",
    "Særlige vilkår: Fargen på boden skal være rød, og taket skal tekkes med "
    "skifer. Høyden må ikke overstige 4,5 meter.",
    "Du kan klage på vedtaket innen tre uker. Klagen sendes til Bergen kommune, "
    "Plan- og bygningsetaten. Ta kontakt dersom du har spørsmål.",
    "Med vennlig hilsen Åse Prøvesen, saksbehandler (oppdiktet).",
]


def make_ocr():
    out = TD / "ocr"
    out.mkdir(parents=True, exist_ok=True)
    body, h1 = styles()
    src = out / "_kilde.pdf"
    doc = SimpleDocTemplate(str(src), pagesize=A4, leftMargin=25 * mm, rightMargin=25 * mm)
    big = ParagraphStyle("big", parent=body, fontSize=13, leading=18)
    doc.build([Paragraph(OCR_TITLE, h1)] + [Paragraph(p, big) for p in OCR_TEXT])
    src2 = out / "_kilde2.pdf"
    doc2 = SimpleDocTemplate(str(src2), pagesize=A4, leftMargin=25 * mm, rightMargin=25 * mm)
    doc2.build([Paragraph("Side med tekstlag", h1),
                Paragraph("Denne siden har allerede et tekstlag og skal hoppes over "
                          "av OCR. Ærlig, søt og åpen.", big)])
    dpi = 200
    img = raster_page(src, 0, dpi)
    images_to_pdf([img], out / "skannet_side.pdf", dpi)
    skew = img.rotate(3, resample=Image.BICUBIC, expand=True, fillcolor=255)
    images_to_pdf([skew], out / "skjev_side.pdf", dpi)
    images_to_pdf([None, img], out / "blandet.pdf", dpi, text_pages={0: (src2, 0)})
    src.unlink()
    src2.unlink()


# ---------------------------------------------------------------------------
# 8. redact
# ---------------------------------------------------------------------------
W1 = [3, 7, 6, 1, 8, 9, 4, 5, 2]
W2 = [5, 4, 3, 2, 7, 6, 5, 4, 3, 2]


def ctrl(digits: list[int], weights: list[int]) -> int | None:
    r = 11 - sum(d * w for d, w in zip(digits, weights)) % 11
    if r == 11:
        return 0
    if r == 10:
        return None
    return r


def fnr_valid(s: str) -> bool:
    d = [int(c) for c in s]
    return len(d) == 11 and ctrl(d[:9], W1) == d[9] and ctrl(d[:10], W2) == d[10]


def make_fnr(ddmmyy: str, start_individ: int) -> tuple[str, dict]:
    for ind in range(start_individ, start_individ + 200):
        base = [int(c) for c in f"{ddmmyy}{ind:03d}"]
        k1 = ctrl(base, W1)
        if k1 is None:
            continue
        k2 = ctrl(base + [k1], W2)
        if k2 is None:
            continue
        s = "".join(map(str, base + [k1, k2]))
        s1 = sum(d * w for d, w in zip(base, W1))
        s2 = sum(d * w for d, w in zip(base + [k1], W2))
        return s, {"individnummer": f"{ind:03d}", "sum1": s1, "k1": k1, "sum2": s2, "k2": k2}
    raise ValueError(ddmmyy)


REDACT = {}


def make_redact():
    out = TD / "redact"
    out.mkdir(parents=True, exist_ok=True)
    f1, c1 = make_fnr("010190", 123)   # written with space
    f2, c2 = make_fnr("150685", 271)
    f3, c3 = make_fnr("290292", 418)
    dn, cd = make_fnr("410388", 152)   # D-number: day + 40
    bad = f2[:10] + str((int(f2[10]) + 1) % 10)
    assert not fnr_valid(bad) and all(fnr_valid(x) for x in (f1, f2, f3, dn))
    f1_spaced = f"{f1[:6]} {f1[6:]}"
    tel1 = "58 01 23 45"
    tel2 = "+47 58 98 76 54"
    email = "kari.fiktivsen@example.no"
    name = "Kari Fiktivsen"
    body, h1 = styles()
    page1 = [
        f"Saken gjelder {name}, født 1. januar 1990, med fødselsnummer {f1_spaced}.",
        f"Hun kan nås på telefon {tel1} eller på e-post {email}.",
        "Søknaden ble mottatt 3. mars 2026 og er registrert med saksnummer "
        "2026/1234. Saksnummeret og datoen skal ikke sladdes.",
        "Beløpet som er søkt om, er 45 000 kroner. Kontonummer er ikke oppgitt.",
    ]
    page2 = [
        f"Medsøker har fødselsnummer {f2}.",
        f"Ektefellen har fødselsnummer {f3}.",
        f"En utenlandsk arbeidstaker i husstanden har D-nummer {dn}.",
        f"Referansenummeret {bad} er et internt løpenummer og ikke et fødselsnummer.",
        f"Kontaktperson hos arbeidsgiver kan nås på {tel2}.",
        f"Vedtaket sendes til {name} med kopi til arkivet.",
    ]
    doc = SimpleDocTemplate(str(out / "personopplysninger.pdf"), pagesize=A4,
                            title="Saksdokument (oppdiktet)", author="Testdata",
                            leftMargin=20 * mm, rightMargin=20 * mm)
    doc.build([Paragraph("Saksdokument – side 1 (oppdiktet)", h1)]
              + [Paragraph(p, body) for p in page1] + [PageBreak()]
              + [Paragraph("Saksdokument – side 2 (oppdiktet)", h1)]
              + [Paragraph(p, body) for p in page2])
    (out / "termer.txt").write_text(name + "\n", encoding="utf-8")
    texts = [norm(t) for t in pdf_text(out / "personopplysninger.pdf")]

    def pages_of(v):
        return [i + 1 for i, t in enumerate(texts) if v in t]

    items = [
        ("fodselsnummer", f1_spaced, f1, c1, True),
        ("fodselsnummer", f2, f2, c2, True),
        ("fodselsnummer", f3, f3, c3, True),
        ("d_nummer", dn, dn, cd, True),
        ("ugyldig_11_siffer", bad, bad, None, False),
        ("telefon", tel1, "58012345", None, True),
        ("telefon", tel2, "+4758987654", None, True),
        ("epost", email, email, None, True),
        ("term", name, name, None, True),
    ]
    REDACT["items"] = []
    for kind, shown, digits, calc, should in items:
        pg = pages_of(shown)
        REDACT["items"].append({"type": kind, "tekst": shown, "normalisert": digits,
                                "sider": pg, "skal_sladdes": should,
                                "beregning": calc})
    (out / "personopplysninger_fasit.json").write_text(
        json.dumps(REDACT, ensure_ascii=False, indent=1), encoding="utf-8")


# ---------------------------------------------------------------------------
# 9. nynorsk
# ---------------------------------------------------------------------------
NB = [
    "Biblioteket i Fiktivhavn holder åpent lenger om kvelden fra neste måned. "
    "Det betyr at flere kan komme innom etter jobb og skole.",
    "Bibliotekaren forteller at de fleste lånerne ønsker seg flere bøker om "
    "friluftsliv, og at en ny samling om fjordglimring blir klar til høsten.",
    "Kommunen har også kjøpt inn en knaskelure, et oppdiktet leseapparat som "
    "skal gjøre det lettere for barn å finne bøker de liker.",
]
NN = [
    "Biblioteket i Fiktivhavn held ope lenger om kvelden frå neste månad. "
    "Det tyder at fleire kan kome innom etter jobb og skule.",
    "Bibliotekaren fortel at dei fleste lånarane ønskjer seg fleire bøker om "
    "friluftsliv, og at ei ny samling om fjordglimring blir klar til hausten.",
    "Kommunen har også kjøpt inn ein knaskelure, eit oppdikta leseapparat som "
    "skal gjere det lettare for born å finne bøker dei likar.",
]


def make_nynorsk():
    out = TD / "nynorsk"
    out.mkdir(parents=True, exist_ok=True)
    (out / "avsnitt_nb.txt").write_text("\n\n".join(NB) + "\n", encoding="utf-8")
    (out / "avsnitt_nn.txt").write_text("\n\n".join(NN) + "\n", encoding="utf-8")
    d = Document()
    d.core_properties.author = "Testdata"
    d.add_heading("Lengre åpningstider på biblioteket", level=1)
    p = d.add_paragraph()
    p.add_run("Biblioteket i Fiktivhavn ").bold = False
    r = p.add_run("holder åpent lenger om kvelden")
    r.bold = True
    p.add_run(" fra neste måned. Det betyr at flere kan komme innom etter jobb og skole.")
    d.add_paragraph(NB[1])
    d.add_paragraph("Nytt tilbud:")
    for item in ["Lengre åpningstid på tirsdager og torsdager",
                 "Egen hylle for fjordglimring",
                 "Lesestund for barn med knaskelure"]:
        d.add_paragraph(item, style="List Bullet")
    d.add_paragraph(NB[2])
    for para in d.paragraphs:
        for run_ in para.runs:
            run_.font.size = Pt(11)
    d.save(out / "avsnitt_nb.docx")


# ---------------------------------------------------------------------------
# Verification
# ---------------------------------------------------------------------------
def verify(sub_data, tables):
    log = []
    ok = lambda cond, msg: (log.append(("OK " if cond else "FEIL ") + msg), cond)[1]
    allok = True

    # compare
    for v in ("v1", "v2"):
        txt = norm(" ".join(pdf_text(TD / "compare" / f"vedtak_{v}.pdf")))
        amount = "1 200 000" if v == "v1" else "1 250 000"
        allok &= ok(amount in txt, f"compare/vedtak_{v}.pdf har tekstlag og {amount}")
        d = Document(TD / "compare" / f"vedtak_{v}.docx")
        allok &= ok(any(amount in p.text for p in d.paragraphs), f"compare/vedtak_{v}.docx har {amount}")
    v2 = norm((TD / "compare" / "vedtak_v2.txt").read_text())
    allok &= ok("Ubrukte midler" not in v2, "compare: avsnittet om ubrukte midler er slettet i v2")

    # metadata
    ex = json.loads(run(["exiftool", "-j", "-n", "-Orientation", "-GPSLatitude", "-GPSLongitude",
                         "-Model", "-Make", "-ImageWidth", "-ImageHeight",
                         str(TD / "metadata" / "bilde_gps.jpg")]).stdout)[0]
    allok &= ok(ex.get("Orientation") == 6 and abs(ex["GPSLatitude"] - BRYGGEN[0]) < 1e-4
                and abs(ex["GPSLongitude"] - BRYGGEN[1]) < 1e-4 and ex.get("Model") == "Fiktiv X100"
                and (ex["ImageWidth"], ex["ImageHeight"]) == (640, 480),
                f"metadata/bilde_gps.jpg EXIF: {ex}")
    with zipfile.ZipFile(TD / "metadata" / "dokument_forfatter.docx") as z:
        core = z.read("docProps/core.xml").decode()
        app = z.read("docProps/app.xml").decode()
    allok &= ok("<dc:creator>Ola Testesen</dc:creator>" in core
                and "<cp:lastModifiedBy>Kari Fiktivsen</cp:lastModifiedBy>" in core
                and "<Company>Eksempelfirma AS</Company>" in app,
                "metadata/dokument_forfatter.docx: dc:creator, cp:lastModifiedBy og Company i docProps")
    with fitz.open(TD / "metadata" / "dokument_forfatter.pdf") as d:
        info = d.xref_get_key(-1, "Info")[1]
        ix = int(info.split()[0])
        allok &= ok(d.metadata["author"] == "Ola Testesen"
                    and "Eksempelfirma" in d.xref_get_key(ix, "Company")[1]
                    and "Kari Fiktivsen" in d.xref_get_key(ix, "LastModifiedBy")[1]
                    and "Ola Testesen" in d.get_xml_metadata(),
                    "metadata/dokument_forfatter.pdf: Info (Author, Company, LastModifiedBy) og XMP")

    # keyframes
    def probe(p, *args):
        return json.loads(run(["ffprobe", "-v", "error", "-of", "json", *args, str(p)]).stdout)
    k = probe(TD / "keyframes" / "klipp.mp4", "-show_format", "-show_streams")
    dur = float(k["format"]["duration"])
    allok &= ok(abs(dur - 10) < 0.1, f"keyframes/klipp.mp4 varighet {dur:.2f} s")
    res = subprocess.run(["ffmpeg", "-hide_banner", "-i", str(TD / "keyframes" / "klipp.mp4"),
                          "-vf", "select='gt(scene,0.3)',showinfo", "-f", "null", "-"],
                         capture_output=True, text=True)
    cuts = [round(float(x), 2) for x in re.findall(r"pts_time:([\d.]+)", res.stderr)]
    LOUD["cuts"] = cuts
    allok &= ok(len(cuts) == 3 and all(abs(c - e) <= 0.2 for c, e in zip(cuts, (2, 5, 8))),
                f"keyframes/klipp.mp4 scenedeteksjon (0.3) gir klipp ved {cuts}")
    s = probe(TD / "keyframes" / "stående.mov", "-show_streams", "-select_streams", "v:0")["streams"][0]
    rot_tag = s.get("tags", {}).get("rotate")
    rot_sd = [sd.get("rotation") for sd in s.get("side_data_list", []) if "rotation" in sd]
    LOUD["rot"] = (s["width"], s["height"], rot_tag, rot_sd)
    allok &= ok(rot_tag == "90" and rot_sd == [-90],
                f"keyframes/stående.mov {s['width']}x{s['height']}, rotate={rot_tag}, displaymatrix={rot_sd}")
    frame = TD / "keyframes" / "_f.png"
    run(["ffmpeg", "-y", "-hide_banner", "-loglevel", "error", "-i",
         str(TD / "keyframes" / "stående.mov"), "-frames:v", "1", str(frame)])
    with Image.open(frame) as fr:
        fsize = fr.size
    frame.unlink()
    LOUD["rot_frame"] = fsize
    allok &= ok(fsize == (180, 320), f"keyframes/stående.mov: ffmpeg (autorotate) gir bilde {fsize[0]}x{fsize[1]}")

    # loudness
    for name, tgt in (("lav.wav", -30), ("høy.wav", -8)):
        i, tp = LOUD[name]
        allok &= ok(abs(i - tgt) <= 0.5, f"loudness/{name}: {i} LUFS, true peak {tp} dBFS")
    a = probe(TD / "loudness" / "video_med_lyd.mp4", "-show_streams")["streams"]
    allok &= ok({x["codec_type"] for x in a} == {"video", "audio"},
                f"loudness/video_med_lyd.mp4 har bilde og lyd, {LOUD['video_med_lyd.mp4']}")

    # subtitles
    segs = sub_data["segments"]
    durs = [s["end"] - s["start"] for s in segs]
    allok &= ok(170 <= segs[-1]["end"] <= 190, f"subtitles: total {segs[-1]['end']:.1f} s, {len(segs)} segmenter")
    allok &= ok(max(durs) > 7.0, f"subtitles: lengste segment {max(durs):.2f} s")
    allok &= ok(min(durs) < 1.2, f"subtitles: korteste segment {min(durs):.2f} s")
    mono = all(w["start"] <= w["end"] and s["start"] <= w["start"] and w["end"] <= s["end"]
               for s in segs for w in s["words"])
    allok &= ok(mono, "subtitles: ordtidskoder ligger innenfor segmentene")
    nw = json.loads((TD / "subtitles" / "transkripsjon_uten_ord.json").read_text())
    allok &= ok(all("words" not in s for s in nw["segments"]), "subtitles: uten_ord mangler words")

    # pdf-tables
    with fitz.open(TD / "pdf-tables" / "budsjett.pdf") as d:
        t = norm(" ".join(p.get_text() for p in d))
        found = [tb for p in d for tb in p.find_tables().tables]
    allok &= ok("1 234 567,50" in t and "(" in t and MINUS in t,
                "pdf-tables/budsjett.pdf har 1 234 567,50, parentes og minustegn i tekstlaget")
    allok &= ok(len(tables["table2_pages"]) >= 2, f"pdf-tables: tabell 2 går over sidene {tables['table2_pages']}")
    LOUD["find_tables"] = [(tb.row_count, tb.col_count) for tb in found]
    with fitz.open(TD / "pdf-tables" / "skannet_tabell.pdf") as d:
        allok &= ok(all(not p.get_text().strip() for p in d), "pdf-tables/skannet_tabell.pdf har ikke tekstlag")

    # ocr
    for fn, pages_text in (("skannet_side.pdf", [False]), ("skjev_side.pdf", [False]),
                           ("blandet.pdf", [True, False])):
        with fitz.open(TD / "ocr" / fn) as d:
            has = [bool(p.get_text().strip()) for p in d]
        allok &= ok(has == pages_text, f"ocr/{fn} tekstlag per side: {has}")
    if shutil.which("tesseract"):
        img = raster_page(TD / "ocr" / "skannet_side.pdf", 0, 200)
        tmp = TD / "ocr" / "_t.png"
        img.save(tmp)
        r = subprocess.run(["tesseract", str(tmp), "-", "-l", "nor"], capture_output=True, text=True)
        tmp.unlink()
        allok &= ok("Bergen kommune" in r.stdout and all(c in r.stdout for c in "æøå"),
                    "ocr/skannet_side.pdf: tesseract leser «Bergen kommune», æ, ø, å")

    # redact
    texts = [norm(t) for t in pdf_text(TD / "redact" / "personopplysninger.pdf")]
    for it in REDACT["items"]:
        allok &= ok(bool(it["sider"]), f"redact: «{it['tekst']}» på side {it['sider']}")
    allok &= ok(REDACT["items"][-1]["sider"] == [1, 2] and sum(t.count("Kari Fiktivsen") for t in texts) == 2,
                "redact: «Kari Fiktivsen» forekommer to ganger")

    # nynorsk
    d = Document(TD / "nynorsk" / "avsnitt_nb.docx")
    styles_ = [p.style.name for p in d.paragraphs]
    allok &= ok("Heading 1" in styles_ and styles_.count("List Bullet") == 3
                and any(r.bold for p in d.paragraphs for r in p.runs),
                "nynorsk/avsnitt_nb.docx har overskrift, fet tekst og punktliste")

    return allok, log


def main():
    if TD.exists():
        for p in TD.iterdir():
            if p.is_dir():
                shutil.rmtree(p)
    TD.mkdir(exist_ok=True)
    for tool in ("ffmpeg", "ffprobe", "exiftool"):
        if not shutil.which(tool):
            raise SystemExit(f"Mangler {tool} i PATH")
    make_compare()
    make_metadata()
    make_keyframes()
    make_loudness()
    sub = make_subtitles()
    tables = make_pdf_tables()
    make_ocr()
    make_redact()
    make_nynorsk()
    # datawrapper: own script (openpyxl and a small .xls writer), see testdata/FASIT.md.
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    import lag_testdata_datawrapper

    lag_testdata_datawrapper.main()
    allok, log = verify(sub, tables)
    print("\n".join(log))
    print(json.dumps({k: v for k, v in LOUD.items()}, ensure_ascii=False, default=str))
    print("Skrift:", FONT_FILE or "Helvetica (innebygd, U+2212 erstattet med -)")
    if not allok:
        raise SystemExit("Verifisering feilet")
    print("Alle testfiler laget og verifisert.")


if __name__ == "__main__":
    main()
